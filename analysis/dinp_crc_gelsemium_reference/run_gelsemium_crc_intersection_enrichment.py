#!/usr/bin/env python3
"""Rebuild the DINP--CRC intersection using the published 1,893-gene set.

The CRC background is Supplementary Table S2 from Que et al., BMC
Complementary Medicine and Therapies (2021), DOI: 10.1186/s12906-021-03273-7.
The DINP side is the frozen 93-gene CTD + ToxCast/Tox21 + T3DB matrix already
used in this project.  The query is the symbol-level intersection of the two
sets.  GO:BP, KEGG, and Reactome are queried separately for two explicit
backgrounds: the 1,893 published CRC genes and the 93 DINP genes.

This runner does not reuse the historical 881-gene GeneCards analysis and does
not use CRC outcome statistics to select the query.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import time
import urllib.error
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl


HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parents[1]
INPUT_DIR = HERE / "inputs"
OUT = HERE / "outputs"
RAW = OUT / "raw_gprofiler"
SUPPLEMENT_XLSX = INPUT_DIR / "12906_2021_3273_MOESM1_ESM.xlsx"
DINP_MATRIX = (
    REPO_ROOT
    / "analysis"
    / "dinp_crc_multi_database_target_convergence"
    / "outputs"
    / "dinp_exposure_gene_matrix_ctd_toxcast_tox21_t3db.csv"
)
API_URL = "https://biit.cs.ut.ee/gprofiler/api/gost/profile/"
SOURCES = ["GO:BP", "KEGG", "REAC"]
USER_AGENT = "whynot17-gelsemium-crc-1893-ora/1.0"
DOI = "10.1186/s12906-021-03273-7"


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def repo_relative(path: Path) -> str:
    return path.resolve().relative_to(REPO_ROOT.resolve()).as_posix()


def clean_symbol(value: Any) -> str:
    return str(value or "").strip().upper()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def load_crc_table() -> list[dict[str, Any]]:
    if not SUPPLEMENT_XLSX.exists():
        raise FileNotFoundError(f"Missing official supplementary file: {SUPPLEMENT_XLSX}")
    workbook = openpyxl.load_workbook(SUPPLEMENT_XLSX, read_only=True, data_only=True)
    sheet_name = " Table S2"
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Expected worksheet {sheet_name!r}; found {workbook.sheetnames}")
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(min_row=3, values_only=True))
    records: list[dict[str, Any]] = []
    for row in rows:
        if row[1] in (None, ""):
            continue
        records.append(
            {
                "article_rank": int(row[0]),
                "gene_symbol_raw": str(row[1]).strip(),
                "gene_symbol": clean_symbol(row[1]),
                "description": row[2] or "",
                "gifts": row[3] if row[3] is not None else "",
                "genecard_id": row[4] or "",
                "relevance_score": row[5] if row[5] is not None else "",
            }
        )
    if len(records) != 1893:
        raise ValueError(f"Expected 1,893 Table S2 rows, found {len(records)}")
    symbols = [record["gene_symbol"] for record in records]
    if len(set(symbols)) != 1893:
        raise ValueError("Table S2 does not contain 1,893 unique normalized symbols")
    ranks = [record["article_rank"] for record in records]
    if ranks != list(range(1, 1894)):
        raise ValueError("Table S2 ranks are not the complete sequence 1..1893")
    return records


def load_dinp_matrix() -> list[dict[str, str]]:
    rows = read_csv(DINP_MATRIX)
    if len(rows) != 93:
        raise ValueError(f"Expected 93 DINP genes, found {len(rows)}")
    symbols = [clean_symbol(row.get("gene_symbol", "")) for row in rows]
    if len(set(symbols)) != 93 or "" in symbols:
        raise ValueError("DINP matrix must contain 93 unique non-empty gene symbols")
    return rows


def log_combination(n: int, k: int) -> float:
    if k < 0 or k > n:
        return float("-inf")
    return math.lgamma(n + 1) - math.lgamma(k + 1) - math.lgamma(n - k + 1)


def hypergeom_sf(population: int, successes: int, draws: int, observed: int) -> float:
    """P[X >= observed] for a hypergeometric random variable."""
    if population <= 0 or successes < 0 or draws < 0 or observed < 0:
        return float("nan")
    if successes > population or draws > population:
        return float("nan")
    if observed == 0:
        return 1.0
    lo = max(observed, draws - (population - successes))
    hi = min(successes, draws)
    if lo > hi:
        return 0.0
    denominator = log_combination(population, draws)
    terms = [
        log_combination(successes, x)
        + log_combination(population - successes, draws - x)
        - denominator
        for x in range(lo, hi + 1)
    ]
    peak = max(terms)
    return min(1.0, math.exp(peak) * sum(math.exp(term - peak) for term in terms))


def request_gprofiler(query: list[str], background: list[str]) -> tuple[dict[str, Any], bytes, dict[str, Any]]:
    payload = {
        "organism": "hsapiens",
        "query": query,
        "background": background,
        "sources": SOURCES,
        "domain_scope": "custom",
        "user_threshold": 1.0,
        "significance_threshold_method": "fdr",
        "all_results": True,
        "no_evidences": False,
        "ordered": False,
        "combined": False,
        "measure_underrepresentation": False,
    }
    body = json.dumps(payload).encode("utf-8")
    request = urllib.request.Request(
        API_URL,
        data=body,
        headers={"Content-Type": "application/json", "User-Agent": USER_AGENT},
        method="POST",
    )
    last_error: Exception | None = None
    for attempt in range(1, 4):
        try:
            with urllib.request.urlopen(request, timeout=300) as response:
                raw = response.read()
            result = json.loads(raw.decode("utf-8"))
            if "result" not in result or "meta" not in result:
                raise ValueError("g:Profiler response lacks result/meta")
            return result, raw, payload
        except (urllib.error.HTTPError, urllib.error.URLError, TimeoutError, ValueError) as exc:
            last_error = exc
            if attempt < 3:
                time.sleep(5 * attempt)
    raise RuntimeError(f"g:Profiler request failed after 3 attempts: {last_error}")


def parse_response(
    result: dict[str, Any], background_name: str, query_symbols: list[str]
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    metadata = result.get("meta", {}).get("result_metadata", {}) or {}
    rows: list[dict[str, Any]] = []
    for item in result.get("result", []):
        if not isinstance(item, dict) or item.get("source") not in SOURCES:
            continue
        domain = int(item.get("effective_domain_size", 0) or 0)
        term_size = int(item.get("term_size", 0) or 0)
        query_size = int(item.get("query_size", 0) or 0)
        overlap = int(item.get("intersection_size", 0) or 0)
        raw_p = hypergeom_sf(domain, term_size, query_size, overlap)
        intersections = item.get("intersections", [])
        if isinstance(intersections, list) and len(intersections) == len(query_symbols):
            # With no_evidences=false, g:Profiler returns one evidence-code
            # list per query item, in the same order as the submitted query.
            # Reconstruct the actual overlapping symbols rather than writing
            # the evidence codes themselves into the result table.
            overlap_genes = [
                symbol
                for symbol, evidence in zip(query_symbols, intersections)
                if evidence
            ]
        else:
            overlap_genes = []
        rows.append(
            {
                "background": background_name,
                "source": item.get("source", ""),
                "term_id": item.get("native", ""),
                "term_name": item.get("name", ""),
                "description": item.get("description", ""),
                "term_size": term_size,
                "effective_domain_size": domain,
                "query_size": query_size,
                "intersection_size": overlap,
                "raw_hypergeom_p": raw_p,
                "gprofiler_adjusted_p": item.get("p_value", ""),
                "gprofiler_significant": item.get("significant", ""),
                "intersection_genes": ";".join(overlap_genes),
                "parents": ";".join(str(value) for value in item.get("parents", []) or []),
            }
        )
    source_sizes = {
        source: int((metadata.get(source) or {}).get("number_of_terms", 0) or 0)
        for source in SOURCES
    }
    audit = {
        "background": background_name,
        "api_rows": len(rows),
        "source_family_sizes": source_sizes,
        "timestamp": result.get("meta", {}).get("timestamp"),
        "gprofiler_version": result.get("meta", {}).get("version"),
        "genes_metadata": result.get("meta", {}).get("genes_metadata", {}),
    }
    return rows, audit


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    RAW.mkdir(parents=True, exist_ok=True)

    crc_rows = load_crc_table()
    dinp_rows = load_dinp_matrix()
    crc_symbols = {row["gene_symbol"] for row in crc_rows}
    dinp_symbols = {clean_symbol(row["gene_symbol"]) for row in dinp_rows}
    intersection = sorted(crc_symbols & dinp_symbols)
    if len(intersection) != 41:
        raise ValueError(f"Expected 41 DINP--CRC intersection genes, found {len(intersection)}")

    dinp_by_symbol = {clean_symbol(row["gene_symbol"]): row for row in dinp_rows}
    crc_out: list[dict[str, Any]] = []
    for row in crc_rows:
        dinp = dinp_by_symbol.get(row["gene_symbol"], {})
        crc_out.append(
            {
                "source_article": DOI,
                "source_table": "Table S2",
                **row,
                "in_dinp_93": row["gene_symbol"] in dinp_symbols,
                "dinp_support_count": dinp.get("exposure_support_count", ""),
                "CTD": dinp.get("CTD", ""),
                "ToxCast": dinp.get("ToxCast", ""),
                "Tox21": dinp.get("Tox21", ""),
                "T3DB": dinp.get("T3DB", ""),
            }
        )
    write_csv(
        OUT / "crc_genecards_1893_normalized.csv",
        crc_out,
        [
            "source_article",
            "source_table",
            "article_rank",
            "gene_symbol_raw",
            "gene_symbol",
            "description",
            "gifts",
            "genecard_id",
            "relevance_score",
            "in_dinp_93",
            "dinp_support_count",
            "CTD",
            "ToxCast",
            "Tox21",
            "T3DB",
        ],
    )

    intersection_rows = []
    for symbol in intersection:
        source = dinp_by_symbol[symbol]
        intersection_rows.append(
            {
                "gene_symbol": symbol,
                "crc_source": "Gelsemium paper Table S2",
                "crc_article_rank": next(row["article_rank"] for row in crc_rows if row["gene_symbol"] == symbol),
                "crc_relevance_score": next(row["relevance_score"] for row in crc_rows if row["gene_symbol"] == symbol),
                **{key: source.get(key, "") for key in ["CTD", "ToxCast", "Tox21", "T3DB", "exposure_support_count"]},
            }
        )
    write_csv(
        OUT / "dinp_crc_intersection_41.csv",
        intersection_rows,
        [
            "gene_symbol",
            "crc_source",
            "crc_article_rank",
            "crc_relevance_score",
            "CTD",
            "ToxCast",
            "Tox21",
            "T3DB",
            "exposure_support_count",
        ],
    )
    write_csv(OUT / "query_41_genes.csv", [{"gene_symbol": gene} for gene in intersection], ["gene_symbol"])
    write_csv(OUT / "background_crc_1893.csv", [{"gene_symbol": gene} for gene in sorted(crc_symbols)], ["gene_symbol"])
    write_csv(OUT / "background_dinp_93.csv", [{"gene_symbol": gene} for gene in sorted(dinp_symbols)], ["gene_symbol"])

    backgrounds = {
        "gelsemium_crc1893": sorted(crc_symbols),
        "dinp93": sorted(dinp_symbols),
    }
    all_enrichment: list[dict[str, Any]] = []
    api_audit: list[dict[str, Any]] = []
    for background_name, background in backgrounds.items():
        raw_path = RAW / f"gprofiler_{background_name}.json"
        payload_path = OUT / f"request_payload_{background_name}.json"
        cached_has_intersections = False
        if raw_path.exists() and payload_path.exists():
            raw = raw_path.read_bytes()
            result = json.loads(raw.decode("utf-8"))
            payload = json.loads(payload_path.read_text(encoding="utf-8"))
            cached_has_intersections = any(
                isinstance(item, dict) and "intersections" in item
                for item in result.get("result", [])
            )
        if not cached_has_intersections:
            result, raw, payload = request_gprofiler(intersection, background)
            raw_path.write_bytes(raw)
            payload_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        rows, audit = parse_response(result, background_name, intersection)
        for row in rows:
            row["query_input_symbols"] = len(intersection)
            row["background_input_symbols"] = len(background)
        all_enrichment.extend(rows)
        api_audit.append(
            {
                **audit,
                "raw_response": repo_relative(raw_path),
                "raw_response_sha256": sha256_file(raw_path),
                "request_payload": repo_relative(payload_path),
                "query_input_symbols": len(intersection),
                "background_input_symbols": len(background),
            }
        )

    fields = [
        "background",
        "source",
        "term_id",
        "term_name",
        "description",
        "term_size",
        "effective_domain_size",
        "query_size",
        "intersection_size",
        "raw_hypergeom_p",
        "gprofiler_adjusted_p",
        "gprofiler_significant",
        "intersection_genes",
        "parents",
        "query_input_symbols",
        "background_input_symbols",
    ]
    write_csv(OUT / "enrichment_all_backgrounds.csv", all_enrichment, fields)
    for source in SOURCES:
        write_csv(
            OUT / f"{source.replace(':', '_').lower()}_all_backgrounds.csv",
            [row for row in all_enrichment if row["source"] == source],
            fields,
        )

    background_summary = {}
    for item in api_audit:
        metadata = item.get("genes_metadata", {}) or {}
        query_meta = metadata.get("query", {}).get("query_1", {}) if isinstance(metadata, dict) else {}
        background_summary[item["background"]] = {
            "input_symbols": item["background_input_symbols"],
            "query_symbols": len(intersection),
            "api_rows": item["api_rows"],
            "gprofiler_version": item.get("gprofiler_version"),
            "query_mapped_ensg": len(query_meta.get("ensgs", []) or []) if isinstance(query_meta, dict) else None,
            "effective_domain_sizes": sorted(
                {row["effective_domain_size"] for row in all_enrichment if row["background"] == item["background"]}
            ),
            "significant_terms_by_source": {
                source: sum(
                    1
                    for row in all_enrichment
                    if row["background"] == item["background"]
                    and row["source"] == source
                    and float(row["gprofiler_adjusted_p"]) < 0.05
                )
                for source in SOURCES
            },
        }

    manifest = {
        "analysis": "DINP--CRC intersection using published Gelsemium CRC GeneCards set",
        "run_utc": utc_now(),
        "paper": {
            "doi": DOI,
            "title": "A network pharmacology-based investigation on the bioactive ingredients and molecular mechanisms of Gelsemium elegans Benth against colorectal cancer",
            "supplementary_table": "Table S2",
            "supplementary_file": repo_relative(SUPPLEMENT_XLSX),
            "supplementary_sha256": sha256_file(SUPPLEMENT_XLSX),
            "reported_crc_candidates": 1893,
        },
        "inputs": {
            "dinp_matrix": repo_relative(DINP_MATRIX),
            "dinp_matrix_sha256": sha256_file(DINP_MATRIX),
            "dinp_input_genes": 93,
            "crc_input_genes": 1893,
            "intersection_genes": 41,
        },
        "gene_normalization": "strip whitespace and uppercase symbols; no related-phthalate substitution or disease-result filtering",
        "gprofiler": {
            "api_url": API_URL,
            "sources": SOURCES,
            "domain_scope": "custom",
            "all_results": True,
            "backgrounds": background_summary,
            "raw_audit": api_audit,
        },
        "outputs": {
            "crc_normalized": repo_relative(OUT / "crc_genecards_1893_normalized.csv"),
            "intersection": repo_relative(OUT / "dinp_crc_intersection_41.csv"),
            "enrichment": repo_relative(OUT / "enrichment_all_backgrounds.csv"),
        },
    }
    (OUT / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    lines = [
        "# DINP--CRC restart using the published 1,893-gene CRC set",
        "",
        "## Frozen inputs",
        "",
        f"- Source: Que et al., BMC Complementary Medicine and Therapies (2021), DOI `{DOI}`.",
        "- CRC source: official Supplementary Table S2, normalized to 1,893 unique gene symbols.",
        "- DINP source: frozen CTD + ToxCast/Tox21 + T3DB matrix, 93 unique gene symbols.",
        f"- Symbol-level intersection: **{len(intersection)} genes**.",
        "- No previous 881-gene GeneCards background or 18-gene query was reused.",
        "",
        "## Enrichment design",
        "",
        "GO:BP, KEGG, and Reactome were queried separately with g:Profiler using `domain_scope=custom` and `all_results=true`.",
        "The 1,893-gene CRC set is the primary source-defined background; the 93-gene DINP set is a sensitivity background.",
        "g:Profiler may report an effective domain smaller or larger than the input symbol count after canonical Ensembl mapping; those values are retained in the output and manifest.",
        "",
        "## Results by background",
        "",
        "| Background | Input genes | Effective domain(s) | GO:BP significant | KEGG significant | Reactome significant |",
        "|---|---:|---|---:|---:|---:|",
    ]
    for name, item in background_summary.items():
        eff = ", ".join(str(value) for value in item["effective_domain_sizes"])
        sig = item["significant_terms_by_source"]
        lines.append(f"| {name} | {item['input_symbols']} | {eff} | {sig['GO:BP']} | {sig['KEGG']} | {sig['REAC']} |")
    lines.extend(
        [
            "",
            "## Files",
            "",
            "- `crc_genecards_1893_normalized.csv`: converted article Table S2.",
            "- `dinp_crc_intersection_41.csv`: 41-gene DINP--CRC intersection with DINP source flags.",
            "- `enrichment_all_backgrounds.csv`: GO:BP/KEGG/Reactome results for both backgrounds.",
            "- `manifest.json`: hashes, API provenance, and effective-domain accounting.",
        ]
    )
    (OUT / "summary.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    print("DINP--CRC GELSEMIUM REFERENCE ANALYSIS: PASS")
    print(f"CRC Table S2 genes: {len(crc_symbols)}")
    print(f"DINP genes: {len(dinp_symbols)}")
    print(f"Intersection genes: {len(intersection)}")
    print(f"Intersection file: {(OUT / 'dinp_crc_intersection_41.csv').resolve()}")
    print(f"Enrichment file: {(OUT / 'enrichment_all_backgrounds.csv').resolve()}")


if __name__ == "__main__":
    main()
