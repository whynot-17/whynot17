from __future__ import annotations

"""GALAXY-led CRC recurrence model, fixed-horizon update, and public validation.

Inputs remain in the project data_raw directory. The script writes aggregate
outputs to --output-root; patient-level predictions are written only under
outputs/local_only and must not be published.
"""

import argparse
import csv
import importlib.util
import json
import math
import re
import statistics
from collections import defaultdict
from pathlib import Path

import numpy as np
import pdfplumber
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

DAY_PER_MONTH = 365.2425 / 12.0
ORIGIN_DAY = 6.0 * DAY_PER_MONTH  # Nominal 6-month landmark (182.62125 days); GALAXY labels 160-200 days as its 6-month timepoint.
HORIZON_DAY = 365.2425
L3_LOW, L3_HIGH, L3_TARGET = 70.0, 112.0, 91.3
L3_TARGET_MONTH = 3.0
MRD0_LOW, MRD0_HIGH, MRD0_TARGET = 14.0, 70.0, 28.0
L6_TARGET_MONTH = 6.0
L6_HALF_QUARTER_MONTH = 1.5
BOOT_REPS = 1000
SEED = 20260930


def clean(x):
    return str(x or "").strip()


def num(x):
    try:
        if x is None or clean(x).lower() in {"", "na", "n/a", "none"}:
            return None
        return float(x)
    except (TypeError, ValueError):
        return None


def status01(x):
    s = clean(x).lower()
    if s in {"positive", "detected", "1"}:
        return 1
    if s in {"negative", "not detected", "0"}:
        return 0
    if isinstance(x, (int, float)) and x in (0, 1):
        return int(x)
    return None


def sex01(x):
    s = clean(x).lower()
    if s in {"male", "m", "man"}:
        return 1
    if s in {"female", "f", "woman"}:
        return 0
    return None


def load_legacy(root: Path):
    path = root / "scripts" / "03_multicohort_landmark_model.py"
    spec = importlib.util.spec_from_file_location("legacy_crc_parser_for_galaxy_model", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def enrich_chen(root: Path, mod):
    patients = mod.read_chen()
    table = Document(root / "data_raw" / "13045_2021_1089_MOESM1_ESM.docx").tables[0]
    headers = [c.text.strip() for c in table.rows[1].cells]
    i_id, i_age, i_sex = headers.index("Patient No."), headers.index("Age"), headers.index("Sex")
    baseline = {}
    for row in table.rows[2:]:
        vals = [c.text.strip() for c in row.cells]
        if vals and vals[i_id]:
            baseline[vals[i_id]] = {"age": num(vals[i_age]), "sex": vals[i_sex]}
    for p in patients:
        p.update(baseline.get(str(p["patient_id"]), {}))
    return patients


def enrich_coloniaiq(root: Path, mod):
    patients = mod.read_coloniaiq()
    reader = PdfReader(root / "data_raw" / "mo_colonaiq_supplement1.pdf")
    baseline = {}
    for page in reader.pages[5:19]:
        for line in (page.extract_text() or "").splitlines():
            tok = line.strip().split()
            if len(tok) >= 6 and re.fullmatch(r"P\d+", tok[0]) and tok[1] in ("Male", "Female"):
                baseline[tok[0]] = {"sex": tok[1], "age": num(tok[2])}
    for p in patients:
        p.update(baseline.get(str(p["patient_id"]), {}))
    return patients


def enrich_cosmos(root: Path, mod):
    patients = mod.read_cosmos()
    ws = load_workbook(root / "data_raw" / "cosmos_crc01_supp_table2.xlsx",
                       read_only=True, data_only=True).active
    baseline = {}
    for r in ws.iter_rows(min_row=4, max_col=23, values_only=True):
        if not r[0]:
            continue
        baseline[str(r[0])] = {
            "age": num(r[1]), "sex": r[2], "site": r[3], "stage": r[4],
            "rfs_months": num(r[13]),
        }
    for p in patients:
        b = baseline.get(str(p["patient_id"]), {})
        p.update(b)
        # Column 14 is RFS from surgery; for non-recurrence patients it is their
        # censor/follow-up time. The legacy parser left this blank for controls.
        if p.get("event_time") is None:
            p["followup_time"] = b.get("rfs_months")
    return patients


def read_galaxy(root: Path):
    path = root / "data_raw" / "galaxy_2024_supplementary_tables_1_3.pdf"
    out = []
    with pdfplumber.open(path) as pdf:
        if len(pdf.pages) != 46:
            raise RuntimeError(f"GALAXY supplement page count changed: {len(pdf.pages)}")
        for page in pdf.pages[2:45]:
            tables = page.extract_tables()
            if not tables:
                continue
            for row in tables[0]:
                if len(row) < 19 or not re.fullmatch(r"\s*\d+\s*", clean(row[0])):
                    continue
                if clean(row[1]) not in {"Male", "Female"}:
                    continue
                r = [clean(x) for x in row[:19]]
                rec = clean(r[4]).lower() == "yes"
                dfs = num(r[6])
                out.append({
                    "cohort": "GALAXY", "patient_id": f"GALAXY-{int(r[0]):04d}",
                    "age": num(r[2]), "sex": r[1], "stage": r[3], "site": None,
                    "event_time_days": dfs if rec else None,
                    "followup_days": None if rec else dfs,
                    "recurred": rec, "dfs_days": dfs,
                    "mrd0": status01(r[10]), "mrd0_day": num(r[11]),
                    "mrd3": status01(r[12]), "mrd6": status01(r[13]),
                    "mrd0_day_exact": True, "mrd3_day_exact": False,
                    "mrd6_day_exact": False,
                })
    ids = [p["patient_id"] for p in out]
    if len(out) != 2240 or len(set(ids)) != 2240:
        raise RuntimeError(f"Expected 2,240 unique GALAXY rows, found {len(out)}")
    return out


def source_patients(root: Path):
    mod = load_legacy(root)
    galaxy = read_galaxy(root)
    cohorts = {
        "GALAXY": galaxy,
        "Chen": enrich_chen(root, mod),
        "ColonAiQ": enrich_coloniaiq(root, mod),
        "COSMOS": enrich_cosmos(root, mod),
    }
    for cohort, patients in cohorts.items():
        for p in patients:
            p["cohort"] = cohort
            p["sex_male"] = sex01(p.get("sex"))
            p["age"] = num(p.get("age"))
            p["stage_high"] = mod.stage_high(p.get("stage"))
            if cohort != "GALAXY":
                p["event_time_days"] = (num(p.get("event_time")) * DAY_PER_MONTH
                                         if num(p.get("event_time")) is not None else None)
                p["followup_days"] = (num(p.get("followup_time")) * DAY_PER_MONTH
                                       if num(p.get("followup_time")) is not None else None)
                hist = []
                for t, s in p.get("history", []):
                    if num(t) is not None and status01(s) is not None and float(t) > 0:
                        hist.append((float(t) * DAY_PER_MONTH, int(status01(s))))
                p["history_days"] = hist
                p["mrd0_status"], p["mrd0_day"] = pick_window(hist, MRD0_LOW, MRD0_HIGH - 1e-6, MRD0_TARGET)
                if cohort == "COSMOS":
                    l6_candidates = [(t, s) for t, s in hist
                                     if abs(t / DAY_PER_MONTH - L6_TARGET_MONTH) < 1e-6]
                else:
                    l6_candidates = [(t, s) for t, s in hist
                                     if abs(t / DAY_PER_MONTH - L6_TARGET_MONTH) <= L6_HALF_QUARTER_MONTH]
                if l6_candidates:
                    p["l6_day"], p["l6_status"] = min(
                        l6_candidates, key=lambda z: (abs(z[0] / DAY_PER_MONTH - L6_TARGET_MONTH), z[0]))
                    p["l6_status"] = int(p["l6_status"])
                    p["l6_day"] = float(p["l6_day"])
                else:
                    p["l6_status"], p["l6_day"] = None, None
                # Use each study's own L3 node. Chen's published schedule starts
                # postoperative ctDNA at month 6, so it has no designated month-3
                # assessment; an isolated sample near month 4 is not relabeled L3.
                if cohort == "COSMOS":
                    l3_candidates = [(t, s) for t, s in hist
                                     if abs(t / DAY_PER_MONTH - L3_TARGET_MONTH) < 1e-6
                                     and p.get("l6_day") is not None and t < p["l6_day"]]
                elif cohort == "ColonAiQ":
                    l3_candidates = [(t, s) for t, s in hist
                                     if abs(t / DAY_PER_MONTH - L3_TARGET_MONTH) <= L6_HALF_QUARTER_MONTH
                                     and p.get("l6_day") is not None and t < p["l6_day"]]
                else:
                    l3_candidates = []
                if cohort == "Chen":
                    p["l3_status"], p["l3_day"] = None, None
                elif l3_candidates:
                    p["l3_day"], p["l3_status"] = min(
                        l3_candidates, key=lambda z: (abs(z[0] / DAY_PER_MONTH - L3_TARGET_MONTH), z[0]))
                    p["l3_status"] = int(p["l3_status"])
                    p["l3_day"] = float(p["l3_day"])
                else:
                    p["l3_status"], p["l3_day"] = None, None
                p["mrd0_day_exact"] = p["mrd3_day_exact"] = p["mrd6_day_exact"] = True
                p["mrd3"] = p["l3_status"]
                p["mrd6"] = p["l6_status"]
                p["mrd0"] = p["mrd0_status"]
            else:
                p["mrd0_status"] = p["mrd0"]
                p["l3_status"] = p["mrd3"]
                p["l6_status"] = p["mrd6"]
                p["l3_day"] = p["l6_day"] = None
                p["history_days"] = []
    return cohorts


def pick_window(hist, low, high, target):
    cand = [(t, s) for t, s in hist if low <= t <= high]
    if not cand:
        return None, None
    t, s = min(cand, key=lambda z: (abs(z[0] - target), z[0]))
    return int(s), float(t)


def outcome_at_origin(p, origin=ORIGIN_DAY):
    et, fu = p.get("event_time_days"), p.get("followup_days")
    if et is not None:
        if et <= origin:
            return False, None
        if et <= HORIZON_DAY:
            return True, 1
        return True, 0
    if fu is not None and fu >= HORIZON_DAY:
        return True, 0
    if fu is not None and fu >= origin:
        return True, None
    return False, None


def transition(a, b):
    names = {(0, 0): "negative_to_negative", (0, 1): "negative_to_positive",
             (1, 0): "positive_to_negative", (1, 1): "positive_to_positive"}
    return names[(int(a), int(b))]


def analysis_records(cohorts):
    result = {c: [] for c in cohorts}
    validation = {c: [] for c in cohorts}
    trajectories = {c: [] for c in cohorts}
    audit = []
    timing = []
    for cohort, patients in cohorts.items():
        counts = defaultdict(int)
        for p in patients:
            counts["total"] += 1
            counts["clinical_complete"] += all(p.get(k) is not None for k in ("age", "sex_male", "stage_high"))
            for key in ("mrd0_status", "l3_status", "l6_status"):
                counts[key] += p.get(key) in (0, 1)
            complete = all(p.get(k) in (0, 1) for k in ("mrd0_status", "l3_status", "l6_status"))
            counts["complete_3timepoints"] += complete
            if cohort == "GALAXY":
                prior_values = [s for s in (p.get("mrd0_status"), p.get("l3_status")) if s in (0, 1)]
                pre_l3_values = [p.get("mrd0_status")] if p.get("mrd0_status") in (0, 1) else []
            else:
                l6_day = p.get("l6_day")
                prior_values = [s for t, s in p.get("history_days", [])
                                if l6_day is not None and t < l6_day and s in (0, 1)]
                l3_day = p.get("l3_day")
                pre_l3_values = [s for t, s in p.get("history_days", [])
                                 if l3_day is not None and t < l3_day and s in (0, 1)]
            counts["prior_mrd_history_before_l6"] += bool(prior_values)
            origin_day = ORIGIN_DAY if cohort == "GALAXY" else p.get("l6_day")
            at_risk, y = outcome_at_origin(p, origin_day if origin_day is not None else ORIGIN_DAY)
            counts["event_free_at_landmark"] += at_risk
            counts["known_outcome_by_12m"] += at_risk and y is not None
            counts["events_from_landmark_to_12m"] += at_risk and y == 1
            counts["controls_through_12m"] += at_risk and y == 0
            if cohort != "GALAXY":
                counts["l6_sample_in_cohort_defined_node"] += p.get("l6_day") is not None
            if cohort == "GALAXY":
                counts["l6_sample_in_cohort_defined_node"] += p.get("l6_status") in (0, 1)
            clinical_ok = all(p.get(k) is not None for k in ("age", "sex_male", "stage_high"))
            if at_risk and clinical_ok and complete:
                traj = dict(p)
                traj["y"] = int(y) if y is not None else None
                traj["ever_prior_l3"] = int(any(s == 1 for s in pre_l3_values))
                traj["ever_prior_l6"] = int(any(s == 1 for s in prior_values))
                traj["transition"] = transition(p["l3_status"], p["l6_status"])
                traj["landmark_day"] = origin_day if origin_day is not None else ORIGIN_DAY
                trajectories[cohort].append(traj)
            if (at_risk and y is not None and clinical_ok and p.get("l6_status") in (0, 1)
                    and prior_values):
                val = dict(p)
                val["y"] = int(y)
                val["ever_prior_l6"] = int(any(s == 1 for s in prior_values))
                val["ever_prior_l3"] = int(any(s == 1 for s in pre_l3_values)) if pre_l3_values else 0
                val["transition"] = (transition(p["l3_status"], p["l6_status"])
                                      if p.get("l3_status") in (0, 1) else "l3_unavailable")
                val["has_pre_l6_history"] = True
                val["landmark_day"] = origin_day if origin_day is not None else ORIGIN_DAY
                validation[cohort].append(val)
                counts["validation_risk_set"] += 1
                counts["validation_events"] += int(y == 1)
            # Strict matched-history risk set used for model fitting and comparison.
            if (at_risk and y is not None and complete and clinical_ok):
                rec = dict(p)
                rec["y"] = int(y)
                rec["transition"] = transition(p["l3_status"], p["l6_status"])
                rec["ever_prior_l3"] = int(any(s == 1 for s in pre_l3_values))
                rec["ever_prior_l6"] = int(any(s == 1 for s in prior_values))
                rec["landmark_day"] = origin_day if origin_day is not None else ORIGIN_DAY
                result[cohort].append(rec)
        audit.append({"cohort": cohort, **counts})
        selected_l3_days = [p["l3_day"] for p in patients if p.get("l3_day") is not None]
        selected_l6_days = [p["l6_day"] for p in patients if p.get("l6_day") is not None]
        if cohort == "GALAXY":
            l3_rule = "published study-defined 70-112-day L3 window; patient-specific date absent"
            l6_rule = "published study-defined 160-200-day L6 window; patient-specific date absent; nominal outcome origin day 182.621"
        elif cohort == "COSMOS":
            l3_rule = "scheduled month-3 status column; exact 3-month node"
            l6_rule = "scheduled month-6 assessment; patient-specific recorded 6-month sample is outcome origin"
        elif cohort == "Chen":
            l3_rule = "not available: published postoperative ctDNA schedule starts at month 6"
            l6_rule = "planned month-6 visit; closest observed test within 4.5-7.5 months; actual date is outcome origin"
        else:
            l3_rule = "quarterly surveillance; closest observed test within 1.5 months of month 3 and before selected L6"
            l6_rule = "quarterly surveillance; closest observed test within 4.5-7.5 months of month 6; actual date is outcome origin"
        timing.append({
            "cohort": cohort,
            "L3_rule": l3_rule,
            "L3_records_in_cohort_defined_node": sum(p.get("l3_status") in (0, 1) for p in patients),
            "L3_selected_median_month": float(np.median(selected_l3_days) / DAY_PER_MONTH) if selected_l3_days else None,
            "L3_selected_min_month": float(min(selected_l3_days) / DAY_PER_MONTH) if selected_l3_days else None,
            "L3_selected_max_month": float(max(selected_l3_days) / DAY_PER_MONTH) if selected_l3_days else None,
            "L6_rule": l6_rule,
            "L6_records_in_cohort_defined_node": counts["l6_sample_in_cohort_defined_node"],
            "L6_selected_median_month": float(np.median(selected_l6_days) / DAY_PER_MONTH) if selected_l6_days else None,
            "L6_selected_min_month": float(min(selected_l6_days) / DAY_PER_MONTH) if selected_l6_days else None,
            "L6_selected_max_month": float(max(selected_l6_days) / DAY_PER_MONTH) if selected_l6_days else None,
            "risk_origin": "nominal day 182.621" if cohort == "GALAXY" else "patient-specific actual L6 assessment date",
            "outcome_window": "after cohort-defined L6 assessment through postoperative month 12",
            "patient_specific_L6_date_available": "no" if cohort == "GALAXY" else "yes",
        })
    return result, validation, trajectories, audit, timing


def l3_l6_outcome_records(cohorts):
    """Cohort-specific L3->L6 status transitions with outcomes after L6 through month 12."""
    result = {c: [] for c in cohorts}
    for cohort, patients in cohorts.items():
        for p in patients:
            if p.get("l3_status") not in (0, 1) or p.get("l6_status") not in (0, 1):
                continue
            origin_day = ORIGIN_DAY if cohort == "GALAXY" else p.get("l6_day")
            if origin_day is None:
                continue
            at_risk, y = outcome_at_origin(p, origin_day)
            if not at_risk or y is None:
                continue
            if cohort == "GALAXY":
                pre_l3 = [s for s in (p.get("mrd0_status"),) if s in (0, 1)]
                pre_l6 = [s for s in (p.get("mrd0_status"), p.get("l3_status")) if s in (0, 1)]
            else:
                hist = p.get("history_days", [])
                l3_day = p.get("l3_day")
                l6_day = p.get("l6_day")
                pre_l3 = [s for t, s in hist if l3_day is not None and t < l3_day and s in (0, 1)]
                pre_l6 = [s for t, s in hist if l6_day is not None and t < l6_day and s in (0, 1)]
            row = dict(p)
            row["y"] = int(y)
            row["transition"] = transition(p["l3_status"], p["l6_status"])
            row["ever_prior_l3"] = int(any(s == 1 for s in pre_l3))
            row["ever_prior_l6"] = int(any(s == 1 for s in pre_l6))
            row["landmark_day"] = origin_day
            result[cohort].append(row)
    return result


def age_parameters(rows):
    a = np.asarray([r["age"] for r in rows], dtype=float)
    mu, sd = float(a.mean()), float(a.std(ddof=0))
    return {"age_mean": mu, "age_sd": sd if sd > 1e-9 else 1.0}


def feature_vector(r, variant, scale):
    base = [float(r["stage_high"]), (float(r["age"]) - scale["age_mean"]) / scale["age_sd"], float(r["sex_male"])]
    names = ["stage_III_IV", "age_z", "sex_male"]
    if variant in ("L3_before", "L6_after"):
        key = "l3_status" if variant == "L3_before" else "l6_status"
        prior = "ever_prior_l3" if variant == "L3_before" else "ever_prior_l6"
        base += [float(r[key]), float(r[prior])]
        names += ["current_mrd_positive_at_3m" if variant == "L3_before" else "current_mrd_positive_at_6m",
                  "ever_prior_mrd_positive"]
    return np.asarray(base, dtype=float), names


def fit_model(rows, variant, mod):
    scale = age_parameters(rows)
    xrows = [feature_vector(r, variant, scale)[0] for r in rows]
    x = np.column_stack([np.ones(len(rows)), np.vstack(xrows)])
    y = np.asarray([r["y"] for r in rows], dtype=int)
    beta = fit_ridge_logit(x, y, penalty=1.0)
    names = ["intercept"] + feature_vector(rows[0], variant, scale)[1]
    return {"variant": variant, "scale": scale, "beta": beta, "terms": names, "ridge_lambda": 1.0}


def predict(model, rows, mod):
    x = np.vstack([feature_vector(r, model["variant"], model["scale"])[0] for r in rows])
    return mod.sigmoid(np.column_stack([np.ones(len(rows)), x]) @ model["beta"])


def metric_values(y, p, mod):
    y = np.asarray(y, dtype=int); p = np.asarray(p, dtype=float)
    lp = np.log(np.clip(p, 1e-8, 1-1e-8) / np.clip(1-p, 1e-8, 1))
    ci, cs = float("nan"), float("nan")
    if len(np.unique(y)) == 2 and float(np.std(lp)) >= 1e-8:
        try:
            cb = fit_ridge_logit(np.column_stack([np.ones(len(lp)), lp]), y, penalty=1e-6)
            if np.all(np.isfinite(cb)) and np.max(np.abs(cb)) < 10:
                ci, cs = float(cb[0]), float(cb[1])
        except (ValueError, np.linalg.LinAlgError, FloatingPointError):
            pass
    return {"auc": mod.rank_auc(y, p), "average_precision": mod.average_precision(y, p),
            "brier": float(np.mean((y-p)**2)), "calibration_intercept": ci,
            "calibration_slope": cs, "mean_predicted_risk": float(np.mean(p))}


def fit_ridge_logit(x, y, penalty=1.0, max_iter=300):
    """Damped Newton solver with an objective line search; intercept is unpenalized."""
    x = np.asarray(x, dtype=float); y = np.asarray(y, dtype=float)
    if len(np.unique(y)) < 2:
        raise ValueError("Both outcome classes are required for logistic fitting")
    npar = x.shape[1]
    beta = np.zeros(npar, dtype=float)
    prevalence = float(np.clip(y.mean(), 1e-5, 1-1e-5))
    beta[0] = math.log(prevalence/(1-prevalence))
    pmat = np.eye(npar, dtype=float) * penalty
    pmat[0, 0] = 0.0

    def objective(b):
        z = x @ b
        return float(np.sum(np.logaddexp(0.0, z) - y*z) + .5*b @ pmat @ b)

    loss = objective(beta)
    for _ in range(max_iter):
        z = np.clip(x @ beta, -35, 35)
        prob = 1.0/(1.0+np.exp(-z))
        w = np.maximum(prob*(1-prob), 1e-9)
        grad = x.T @ (prob-y) + pmat @ beta
        hess = (x.T*w) @ x + pmat + np.eye(npar)*1e-8
        step = np.linalg.solve(hess, grad)
        if float(np.max(np.abs(grad))) < 1e-7:
            break
        alpha = 1.0
        accepted = False
        descent = float(grad @ step)
        while alpha >= 1e-10:
            candidate = beta - alpha*step
            new_loss = objective(candidate)
            if new_loss <= loss - 1e-4*alpha*descent:
                beta, loss, accepted = candidate, new_loss, True
                break
            alpha *= .5
        if not accepted or float(np.max(np.abs(alpha*step))) < 1e-8:
            break
    if not np.all(np.isfinite(beta)):
        raise FloatingPointError("Non-finite ridge logistic coefficients")
    return beta


def percentile_ci(values):
    a = np.asarray([x for x in values if np.isfinite(x)], dtype=float)
    if not len(a):
        return None, None
    return float(np.quantile(a, .025)), float(np.quantile(a, .975))


def bootstrap_internal(rows, mod, reps=BOOT_REPS, seed=SEED):
    rng = np.random.default_rng(seed)
    y = np.asarray([r["y"] for r in rows], dtype=int)
    full = {v: fit_model(rows, v, mod) for v in ("clinical", "L3_before", "L6_after")}
    apparent = {v: metric_values(y, predict(full[v], rows, mod), mod) for v in ("clinical", "L6_after")}
    optimism = {v: defaultdict(list) for v in ("clinical", "L6_after")}
    boot_test = {v: defaultdict(list) for v in ("clinical", "L6_after")}
    oob_sum = {v: np.zeros(len(rows), dtype=float) for v in ("clinical", "L3_before", "L6_after")}
    oob_n = np.zeros(len(rows), dtype=int)
    successes = 0
    for b in range(reps):
        idx = rng.integers(0, len(rows), len(rows))
        boot = [rows[int(i)] for i in idx]
        unique = np.zeros(len(rows), dtype=bool); unique[np.unique(idx)] = True
        try:
            models = {v: fit_model(boot, v, mod) for v in ("clinical", "L3_before", "L6_after")}
            yb = np.asarray([r["y"] for r in boot], dtype=int)
            for v in ("clinical", "L6_after"):
                mb = metric_values(yb, predict(models[v], boot, mod), mod)
                mt = metric_values(y, predict(models[v], rows, mod), mod)
                for metric in ("auc", "average_precision", "brier", "calibration_intercept", "calibration_slope"):
                    if np.isfinite(mb[metric]) and np.isfinite(mt[metric]):
                        optimism[v][metric].append(mb[metric] - mt[metric])
                        boot_test[v][metric].append(mt[metric])
            oob = np.flatnonzero(~unique)
            if len(oob):
                for v in ("clinical", "L3_before", "L6_after"):
                    oob_sum[v][oob] += predict(models[v], [rows[int(i)] for i in oob], mod)
                oob_n[oob] += 1
            successes += 1
        except (ValueError, np.linalg.LinAlgError, FloatingPointError):
            continue
    if successes < reps * .95:
        raise RuntimeError(f"Only {successes}/{reps} bootstrap fits succeeded")
    oob_pred = {v: np.divide(oob_sum[v], oob_n, out=np.full(len(rows), np.nan), where=oob_n > 0)
                for v in ("clinical", "L3_before", "L6_after")}
    rows_out = []
    for v in ("clinical", "L6_after"):
        for metric, app in apparent[v].items():
            if metric not in optimism[v] or not optimism[v][metric]:
                continue
            opt = float(np.mean(optimism[v][metric]))
            corrected_dist = [app - x for x in optimism[v][metric]]
            lo, hi = percentile_ci(corrected_dist)
            rows_out.append({"model": "Clinical-only" if v == "clinical" else "Clinical+MRD",
                             "metric": metric, "apparent": app, "mean_bootstrap_optimism": opt,
                             "optimism_corrected": app - opt, "corrected_ci_low": lo,
                             "corrected_ci_high": hi, "bootstrap_refits": successes})
    return full, rows_out, oob_pred, oob_n, apparent


def calibration_bins(y, p, cohort, model, bins=10):
    order = np.argsort(p)
    chunks = np.array_split(order, min(bins, len(order)))
    out = []
    for i, ids in enumerate(chunks, 1):
        if not len(ids):
            continue
        out.append({"cohort": cohort, "model": model, "bin": i, "n": len(ids),
                    "events": int(np.asarray(y)[ids].sum()),
                    "mean_predicted_risk": float(np.mean(np.asarray(p)[ids])),
                    "observed_event_rate": float(np.mean(np.asarray(y)[ids]))})
    return out


def wilson(x, n):
    if not n:
        return None, None
    z = 1.959963984540054
    p = x / n
    den = 1 + z*z/n
    mid = (p + z*z/(2*n)) / den
    half = z * math.sqrt(p*(1-p)/n + z*z/(4*n*n)) / den
    return max(0.0, mid-half), min(1.0, mid+half)


def bootstrap_group_medians(group, reps=BOOT_REPS, seed=SEED):
    rng = np.random.default_rng(seed + len(group))
    before = np.asarray([r["risk_before_oob"] for r in group], dtype=float)
    after = np.asarray([r["risk_after_oob"] for r in group], dtype=float)
    vals = np.empty((reps, 3), dtype=float)
    for b in range(reps):
        ix = rng.integers(0, len(group), len(group))
        vals[b] = [np.median(before[ix]), np.median(after[ix]), np.median(after[ix]-before[ix])]
    cis = [percentile_ci(vals[:, j]) for j in range(3)]
    return cis


def paired_delta_ci(y, p0, p1, mod, reps=BOOT_REPS, seed=SEED):
    rng = np.random.default_rng(seed + len(y))
    y = np.asarray(y, dtype=int); p0 = np.asarray(p0); p1 = np.asarray(p1)
    store = defaultdict(list); valid = 0
    for _ in range(reps):
        ix = rng.integers(0, len(y), len(y))
        if len(np.unique(y[ix])) < 2:
            continue
        a = metric_values(y[ix], p0[ix], mod); b = metric_values(y[ix], p1[ix], mod)
        for m in ("auc", "average_precision", "brier"):
            if np.isfinite(a[m]) and np.isfinite(b[m]):
                store[m].append(b[m]-a[m])
        valid += 1
    return {m: (*percentile_ci(store[m]), valid) for m in ("auc", "average_precision", "brier")}


def csv_write(path, rows):
    rows = list(rows)
    if not rows:
        return
    keys = list(dict.fromkeys(k for r in rows for k in r))
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=keys, extrasaction="ignore")
        w.writeheader(); w.writerows(rows)


def svg_write(path, body, width=1100, height=700):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}"><rect width="100%" height="100%" fill="white"/>{body}</svg>', encoding="utf-8")


def esc(x):
    return (str(x).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def make_figures(figdir, metrics, updates, memory, transitions, cases):
    figdir.mkdir(parents=True, exist_ok=True)
    body = '<text x="550" y="45" text-anchor="middle" font-size="26" font-family="Arial">GALAXY 主开发与冻结模型公共验证</text>'
    labels = ["GALAXY development", "Freeze coefficients", "Chen validation", "COSMOS validation", "ColonAiQ if compatible", "Future institutional external validation"]
    xs = [90, 440, 790, 790, 790, 1010]
    ys = [115, 115, 70, 160, 250, 115]
    for i, (lab, x, y) in enumerate(zip(labels, xs, ys)):
        body += f'<rect x="{x}" y="{y}" width="220" height="64" rx="12" fill="{("#dbeafe" if i<2 else "#e7f5ec")}" stroke="#334155"/><text x="{x+110}" y="{y+27}" text-anchor="middle" font-size="15" font-family="Arial">{esc(lab)}</text>'
    for y in [102, 192, 282]:
        body += f'<path d="M 1010 147 C 1050 {y}, 1040 {y}, 1010 {y}" fill="none" stroke="#64748b" stroke-width="2"/>'
    body += '<path d="M 310 147 L 430 147" stroke="#334155" stroke-width="3" marker-end="url(#a)"/><defs><marker id="a" markerWidth="10" markerHeight="10" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#334155"/></marker></defs>'
    svg_write(figdir / "study_design.svg", body, 1260, 390)

    cohorts = [r["cohort"] for r in metrics if r["model"] == "Clinical-only"]
    body = '<text x="550" y="35" text-anchor="middle" font-size="23" font-family="Arial">冻结模型的外部区分度、校准与误差</text>'
    panels = [("auc", "ROC-AUC", .5, 1), ("brier", "Brier score", 0, max(.25, max([float(r['brier']) for r in metrics if r.get('brier') not in (None, '')] or [.25]))), ("calibration_slope", "Calibration slope", 0, 2)]
    for pi, (metric, title, lo, hi) in enumerate(panels):
        ox=70+pi*355; oy=95; pw=300; ph=420
        body += f'<text x="{ox+pw/2}" y="{oy-15}" text-anchor="middle" font-size="17" font-family="Arial">{title}</text><line x1="{ox}" y1="{oy+ph}" x2="{ox+pw}" y2="{oy+ph}" stroke="#64748b"/><line x1="{ox}" y1="{oy}" x2="{ox}" y2="{oy+ph}" stroke="#64748b"/>'
        for ci, cohort in enumerate(cohorts):
            barw=26; gap=10; groupw=68; gx=ox+25+ci*groupw
            for mi, model in enumerate(("Clinical-only", "Clinical+MRD")):
                row=next((z for z in metrics if z["cohort"]==cohort and z["model"]==model),None)
                val=float(row[metric]) if row and row.get(metric) not in (None, "") and np.isfinite(float(row[metric])) else None
                if val is not None:
                    vv=min(hi,max(lo,val)); bh=(vv-lo)/(hi-lo)*ph
                    body += f'<rect x="{gx+mi*(barw+4)}" y="{oy+ph-bh}" width="{barw}" height="{bh}" fill="{("#64748b" if mi==0 else "#e07a5f")}"/><text x="{gx+mi*(barw+4)+barw/2}" y="{oy+ph+18}" text-anchor="middle" font-size="10" font-family="Arial">{model.split("-")[0]}</text>'
            body += f'<text x="{gx+30}" y="{oy+ph+42}" text-anchor="middle" font-size="11" font-family="Arial">{esc(cohort)}</text>'
        body += f'<text x="{ox-15}" y="{oy+ph/2}" transform="rotate(-90 {ox-15} {oy+ph/2})" text-anchor="middle" font-size="12" font-family="Arial">Metric</text>'
    svg_write(figdir / "external_validation_auc_brier.svg", body)

    body = '<text x="550" y="38" text-anchor="middle" font-size="22" font-family="Arial">同一(182.6天,12月]结局：L3信息到L6信息的风险更新（OOB）</text>'
    order=["negative_to_negative","negative_to_positive","positive_to_negative","positive_to_positive"]
    zh={"negative_to_negative":"阴→阴","negative_to_positive":"阴→阳","positive_to_negative":"阳→阴","positive_to_positive":"阳→阳"}
    cols={"negative_to_negative":"#238b67","negative_to_positive":"#e67e22","positive_to_negative":"#4c91c6","positive_to_positive":"#b82e38"}
    x0=170; y0=100; rowh=115; scale=750
    for i,state in enumerate(order):
        row=next((r for r in updates if r["transition"]==state),None)
        if not row: continue
        y=y0+i*rowh
        before=float(row["median_risk_before"]); after=float(row["median_risk_after"])
        body += f'<text x="45" y="{y+28}" font-size="16" font-family="Arial">{zh[state]} (n={row["n"]})</text><line x1="{x0}" y1="{y+24}" x2="{x0+before*scale}" y2="{y+24}" stroke="#64748b" stroke-width="7"/><line x1="{x0}" y1="{y+52}" x2="{x0+after*scale}" y2="{y+52}" stroke="{cols[state]}" stroke-width="7"/><text x="{x0+before*scale+8}" y="{y+28}" font-size="12" font-family="Arial">before {before:.1%}</text><text x="{x0+after*scale+8}" y="{y+56}" font-size="12" font-family="Arial">after {after:.1%}; Δ {float(row["median_delta_risk"]):+.1%}</text>'
    body += '<text x="75" y="82" font-size="12" font-family="Arial">灰色：L3可用信息；彩色：L6更新后；均针对同一预测窗口</text>'
    svg_write(figdir / "same_horizon_risk_update.svg", body)

    body = '<text x="550" y="38" text-anchor="middle" font-size="22" font-family="Arial">L6 MRD 阴性患者的分子记忆：从未阳性 vs 既往阳性</text>'
    memcohorts=sorted(set(r["cohort"] for r in memory))
    ybase=105
    for ci,c in enumerate(memcohorts):
        y=ybase+ci*115
        body += f'<text x="40" y="{y+22}" font-size="15" font-family="Arial">{esc(c)}</text>'
        for gi,g in enumerate(("never_positive","previously_positive_now_negative")):
            r=next((x for x in memory if x["cohort"]==c and x["memory_group"]==g),None)
            if r and int(r["n"]):
                x=215+gi*390
                obs=float(r["observed_recurrence_rate"]); pred=float(r["mean_predicted_risk"])
                body += f'<text x="{x}" y="{y+18}" font-size="12" font-family="Arial">{("Never positive" if gi==0 else "Prior +, current -")} n={r["n"]}</text><rect x="{x}" y="{y+30}" width="{obs*230}" height="18" fill="#d95d39"/><text x="{x+240}" y="{y+44}" font-size="12" font-family="Arial">actual {obs:.1%}</text><rect x="{x}" y="{y+58}" width="{pred*230}" height="18" fill="#457b9d"/><text x="{x+240}" y="{y+72}" font-size="12" font-family="Arial">predicted {pred:.1%}</text>'
    svg_write(figdir / "molecular_memory_validation.svg", body, 1150, max(360, ybase+len(memcohorts)*115))

    body = '<text x="650" y="35" text-anchor="middle" font-size="23" font-family="Arial">各队列3月→6月MRD状态转移与(182.6天,12月]复发</text>'
    coh=sorted(set(r["cohort"] for r in transitions))
    trans_order=order
    for ci,c in enumerate(coh):
        y=85+ci*125
        body += f'<text x="40" y="{y+18}" font-size="15" font-family="Arial">{esc(c)}</text>'
        for ti,t in enumerate(trans_order):
            r=next((x for x in transitions if x["cohort"]==c and x["transition"]==t),None)
            if not r or not int(r["n"]): continue
            x=180+ti*245
            rate=float(r["event_rate"])
            body += f'<text x="{x}" y="{y+17}" font-size="12" font-family="Arial">{zh[t]} n={r["n"]}</text><rect x="{x}" y="{y+28}" width="200" height="18" fill="#e5e7eb"/><rect x="{x}" y="{y+28}" width="{200*rate}" height="18" fill="{cols[t]}"/><text x="{x}" y="{y+66}" font-size="12" font-family="Arial">复发 {rate:.1%} ({r["events"]})</text>'
    svg_write(figdir / "dynamic_state_transition.svg", body, 1200, 85+len(coh)*125)

    # Local-only examples are actual records selected from the observed GALAXY cohort.
    body = '<text x="500" y="38" text-anchor="middle" font-size="22" font-family="Arial">GALAXY真实患者的代表性MRD风险轨迹（仅本地）</text>'
    for i,c in enumerate(cases):
        y=105+i*115
        body += f'<text x="45" y="{y+15}" font-size="14" font-family="Arial">{esc(c["transition_zh"])} | {esc(c["patient_id"])} | {esc(c["outcome_label"])}</text><line x1="100" y1="{y+44}" x2="{100+float(c["risk_before_update"])*750}" y2="{y+44}" stroke="#64748b" stroke-width="7"/><line x1="100" y1="{y+72}" x2="{100+float(c["risk_after_update"])*750}" y2="{y+72}" stroke="{c["color"]}" stroke-width="7"/><text x="880" y="{y+49}" font-size="12" font-family="Arial">{float(c["risk_before_update"]):.1%} → {float(c["risk_after_update"]):.1%}</text>'
    svg_write(Path(cases[0]["figure_path"]), body, 1020, 105+len(cases)*115)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--project-root", default=r"E:\crc_mrd_digital_twin")
    ap.add_argument("--output-root", required=True)
    ap.add_argument("--bootstrap", type=int, default=BOOT_REPS)
    args = ap.parse_args()
    root, outroot = Path(args.project_root), Path(args.output_root)
    tables, reports, figs, local, modelsdir = [outroot / d for d in
        ("outputs/tables", "reports", "outputs/figures", "outputs/local_only", "outputs/models")]
    for d in (tables, reports, figs, local, modelsdir): d.mkdir(parents=True, exist_ok=True)

    cohorts = source_patients(root)
    rows_by_cohort, validation_by_cohort, trajectories_by_cohort, audit, timing = analysis_records(cohorts)
    l3_l6_by_cohort = l3_l6_outcome_records(cohorts)
    galaxy = rows_by_cohort["GALAXY"]
    if len(galaxy) < 100 or sum(r["y"] for r in galaxy) < 20:
        raise RuntimeError(f"Insufficient GALAXY L6-window training sample: n={len(galaxy)} events={sum(r['y'] for r in galaxy)}")
    mod = load_legacy(root)
    full_models, boot_rows, oob, oob_n, apparent = bootstrap_internal(galaxy, mod, args.bootstrap)

    # The frozen final models are fitted once on all eligible GALAXY development rows.
    # External cohorts do not enter this step in any way.
    beta_rows = []
    for variant, model in full_models.items():
        model_name = {"clinical":"Clinical-only", "L3_before":"Same-horizon before update", "L6_after":"Clinical+MRD / after update"}[variant]
        for term, coefficient in zip(model["terms"], model["beta"]):
            beta_rows.append({"model": model_name, "term": term, "coefficient": float(coefficient),
                              "odds_ratio_per_unit": float(math.exp(np.clip(coefficient, -30, 30))),
                              **model["scale"], "ridge_lambda": model["ridge_lambda"], "development_n": len(galaxy),
                              "development_events": int(sum(r["y"] for r in galaxy))})
    csv_write(tables / "galaxy_model_coefficients.csv", beta_rows)
    (modelsdir / "galaxy_frozen_ridge_models.json").write_text(json.dumps({
        "development_cohort":"GALAXY", "origin_day":ORIGIN_DAY, "horizon_day":HORIZON_DAY,
        "outcome":"clinical recurrence in (182.62125 days, 365.2425 days]", "ridge_lambda":1.0,
        "models":{k:{"terms":v["terms"],"coefficients":[float(x) for x in v["beta"]],"age_scale":v["scale"]} for k,v in full_models.items()},
        "feature_note":"GALAXY patient-level site is unavailable; clinical core uses stage, age, sex. MRD is binary/assay-agnostic. Ever-prior-positive summarizes observed pre-L6 statuses; prior testing density differs by cohort."
    }, ensure_ascii=False, indent=2), encoding="utf-8")

    boot_table = []
    for r in boot_rows:
        boot_table.append(r)
    # OOB predictions yield the same-horizon update estimates and development calibration curve.
    for i, r in enumerate(galaxy):
        r["risk_before_oob"] = float(oob["L3_before"][i])
        r["risk_after_oob"] = float(oob["L6_after"][i])
    trajectory_labels = {
        "persistent_negative": (0, 0, "持续阴性"),
        "conversion": (1, 0, "转阳"),
        "clearance": (0, 1, "清除"),
        "persistent_positive": (1, 1, "持续阳性"),
    }
    trajectory_main_rows = []
    for name, (current, prior_positive, label) in trajectory_labels.items():
        g = [r for r in galaxy if r["l6_status"] == current and r["ever_prior_l6"] == prior_positive]
        if not g:
            continue
        events = int(sum(r["y"] for r in g))
        ci_low, ci_high = wilson(events, len(g))
        predicted = [r["risk_after_oob"] for r in g if np.isfinite(r["risk_after_oob"])]
        trajectory_main_rows.append({
            "trajectory": name, "trajectory_zh": label,
            "current_l6_status": "positive" if current else "negative",
            "prior_mrd_history": "既往至少一次阳性" if prior_positive else "既往未观察到阳性",
            "n": len(g), "events_6_to_12m": events,
            "recurrence_rate_6_to_12m": events / len(g),
            "recurrence_rate_wilson_95ci_low": ci_low,
            "recurrence_rate_wilson_95ci_high": ci_high,
            "n_with_oob_predicted_risk": len(predicted),
            "mean_oob_predicted_risk": float(np.mean(predicted)) if predicted else None,
            "median_oob_predicted_risk": float(np.median(predicted)) if predicted else None,
            "risk_origin_day": ORIGIN_DAY, "outcome_window": "(182.62125 days, 365.2425 days]",
            "prediction_source": "paired patient-level bootstrap out-of-bag",
        })
    calibration_rows = []
    for key,label in (("clinical","Clinical-only OOB"),
                      ("L3_before","Same-horizon L3 history OOB"),
                      ("L6_after","Clinical+MRD OOB")):
        valid=np.isfinite(oob[key])
        calibration_rows.extend(calibration_bins(np.asarray([r["y"] for r,v in zip(galaxy,valid) if v]),
                                                  oob[key][valid], "GALAXY", label))

    # External validation: apply frozen GALAXY coefficients with no refitting or recalibration.
    metric_rows, increment_rows, calibration_ext = [], [], []
    evaluated = {"GALAXY": galaxy}
    for cohort in ("Chen", "COSMOS", "ColonAiQ"):
        evaluated[cohort] = validation_by_cohort.get(cohort, [])
    all_ext_y, all_ext_c, all_ext_m = [], [], []
    ext_for_memory = {}
    for cohort, rs in evaluated.items():
        if not rs:
            continue
        y=np.asarray([r["y"] for r in rs],dtype=int)
        origin_months=[float(r.get("landmark_day",ORIGIN_DAY))/DAY_PER_MONTH for r in rs]
        pc=predict(full_models["clinical"],rs,mod)
        pm=predict(full_models["L6_after"],rs,mod)
        ext_for_memory[cohort]=(rs,pm)
        for name,pred in (("Clinical-only",pc),("Clinical+MRD",pm)):
            met=metric_values(y,pred,mod)
            metric_rows.append({"cohort":cohort,"model":name,"n":len(rs),"events":int(y.sum()),
                                "event_rate":float(y.mean()),**met,"model_frozen":cohort!="GALAXY",
                                "L6_origin_median_month":float(np.median(origin_months)),
                                "L6_origin_min_month":float(min(origin_months)),"L6_origin_max_month":float(max(origin_months)),
                                "outcome_window":"(day 182.621, postoperative month 12]" if cohort=="GALAXY" else "after individual cohort-defined L6 assessment through postoperative month 12",
                                "validation_role":"development bootstrap" if cohort=="GALAXY" else
                                    ("priority independent validation" if cohort in ("COSMOS", "Chen") else "exploratory; sparse events")})
            calibration_ext.extend(calibration_bins(y,pred,cohort,name))
        deltas=paired_delta_ci(y,pc,pm,mod,BOOT_REPS,SEED+len(cohort))
        for m,(lo,hi,nvalid) in deltas.items():
            inc={"cohort":cohort,"n":len(rs),"events":int(y.sum()),"metric":m,
                 "clinical_only":metric_values(y,pc,mod)[m],"clinical_plus_mrd":metric_values(y,pm,mod)[m],
                 "delta_mrd_minus_clinical":metric_values(y,pm,mod)[m]-metric_values(y,pc,mod)[m],
                 "bootstrap_ci_low":lo,"bootstrap_ci_high":hi,"valid_bootstrap_replicates":nvalid}
            increment_rows.append(inc)
        if cohort in ("COSMOS", "Chen"):
            all_ext_y.extend(y.tolist()); all_ext_c.extend(pc.tolist()); all_ext_m.extend(pm.tolist())
    if all_ext_y:
        y=np.asarray(all_ext_y,dtype=int); pc=np.asarray(all_ext_c); pm=np.asarray(all_ext_m)
        for name,pred in (("Clinical-only",pc),("Clinical+MRD",pm)):
            metric_rows.append({"cohort":"Pooled priority external","model":name,"n":len(y),"events":int(y.sum()),
                                "event_rate":float(y.mean()),**metric_values(y,pred,mod),"model_frozen":True,
                                "L6_origin_median_month":None,"L6_origin_min_month":None,"L6_origin_max_month":None,
                                "outcome_window":"cohort-specific L6 assessment through postoperative month 12",
                                "validation_role":"pooled priority external validation (COSMOS + Chen)"})
        d=paired_delta_ci(y,pc,pm,mod)
        for m,(lo,hi,nvalid) in d.items():
            increment_rows.append({"cohort":"Pooled priority external","n":len(y),"events":int(y.sum()),"metric":m,
                "clinical_only":metric_values(y,pc,mod)[m],"clinical_plus_mrd":metric_values(y,pm,mod)[m],
                "delta_mrd_minus_clinical":metric_values(y,pm,mod)[m]-metric_values(y,pc,mod)[m],
                "bootstrap_ci_low":lo,"bootstrap_ci_high":hi,"valid_bootstrap_replicates":nvalid})
    calibration_rows.extend(calibration_ext)

    # Overall same-horizon performance uses paired OOB predictions, not apparent fits.
    same_y=np.asarray([r["y"] for r in galaxy],dtype=int)
    before_metrics=metric_values(same_y,oob["L3_before"],mod)
    after_metrics=metric_values(same_y,oob["L6_after"],mod)
    same_delta_ci=paired_delta_ci(same_y,oob["L3_before"],oob["L6_after"],mod,BOOT_REPS,SEED+777)
    same_perf_rows=[]
    for metric in ("auc","average_precision","brier"):
        lo,hi,nvalid=same_delta_ci[metric]
        same_perf_rows.append({"metric":metric,"before_L3_history_oob":before_metrics[metric],
            "after_L6_history_oob":after_metrics[metric],
            "delta_after_minus_before":after_metrics[metric]-before_metrics[metric],
            "delta_ci_low":lo,"delta_ci_high":hi,"valid_bootstrap_replicates":nvalid,
            "n":len(galaxy),"events":int(same_y.sum()),"same_outcome_window":"(182.62125 days, 365.2425 days]"})

    # Keep the previously completed day-200 analysis as a secondary sensitivity
    # and make the comparison explicit beside the nominal day-182.6 primary run.
    old_tables = root / "analysis_runs" / "galaxy_2024_l6_l12" / "outputs" / "tables"
    old_boot_path = old_tables / "galaxy_internal_bootstrap_metrics.csv"
    old_external_path = old_tables / "external_validation_metrics.csv"
    old_same_path = old_tables / "same_horizon_model_performance.csv"
    day200_boot, day200_internal = {}, None
    if old_boot_path.exists() and old_external_path.exists():
        with old_boot_path.open(encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f):
                day200_boot[(row["model"], row["metric"])] = row
        with old_external_path.open(encoding="utf-8-sig", newline="") as f:
            day200_internal = next((row for row in csv.DictReader(f)
                                    if row.get("cohort") == "GALAXY" and row.get("model") == "Clinical-only"), None)
    landmark_sensitivity_rows=[]
    landmark_specs=[("nominal_L6_day182.6_primary",len(galaxy),int(sum(r["y"] for r in galaxy)),boot_rows)]
    if day200_internal and day200_boot:
        landmark_specs.append(("day200_secondary_sensitivity",int(day200_internal["n"]),int(day200_internal["events"]),day200_boot))
    for landmark,n_landmark,event_count,source_rows in landmark_specs:
        for model_name in ("Clinical-only","Clinical+MRD"):
            for metric in ("auc","average_precision","brier"):
                q=next(r for r in source_rows if r["model"]==model_name and r["metric"]==metric) if isinstance(source_rows,list) else source_rows[(model_name,metric)]
                landmark_sensitivity_rows.append({"landmark_analysis":landmark,"n":n_landmark,"events":event_count,
                    "model":model_name,"metric":metric,"apparent":float(q["apparent"]),
                    "optimism_corrected":float(q["optimism_corrected"]),
                    "corrected_ci_low":float(q["corrected_ci_low"]),"corrected_ci_high":float(q["corrected_ci_high"]),
                    "bootstrap_refits":int(q.get("bootstrap_refits",args.bootstrap))})
    day200_same={}
    if old_same_path.exists():
        with old_same_path.open(encoding="utf-8-sig", newline="") as f:
            day200_same={row["metric"]:row for row in csv.DictReader(f)}
    same_horizon_sensitivity_rows=[]
    for r in same_perf_rows:
        old=day200_same.get(r["metric"])
        same_horizon_sensitivity_rows.append({"metric":r["metric"],"day1826_n":r["n"],"day1826_events":r["events"],
            "day1826_before_L3_oob":r["before_L3_history_oob"],"day1826_after_L6_oob":r["after_L6_history_oob"],
            "day1826_delta":r["delta_after_minus_before"],
            "day200_n":int(old["n"]) if old else None,"day200_events":int(old["events"]) if old else None,
            "day200_before_L3_oob":float(old["before_L3_history_oob"]) if old else None,
            "day200_after_L6_oob":float(old["after_L6_history_oob"]) if old else None,
            "day200_delta":float(old["delta_after_minus_before"]) if old else None})

    # Cohort-native L3->L6 transitions and L6-to-month-12 outcomes.
    l3l6_audit_rows=[]
    for cohort,patients in cohorts.items():
        both=[p for p in patients if p.get("l3_status") in (0,1) and p.get("l6_status") in (0,1)]
        distinct=[p for p in both if cohort=="GALAXY" or
                  (p.get("l3_day") is not None and p.get("l6_day") is not None and p["l3_day"]<p["l6_day"])]
        rs=l3_l6_by_cohort.get(cohort,[])
        l3l6_audit_rows.append({"cohort":cohort,"L3_assessments":sum(p.get("l3_status") in (0,1) for p in patients),
            "L6_assessments":sum(p.get("l6_status") in (0,1) for p in patients),
            "patients_with_both_statuses":len(both),"patients_with_ordered_L3_before_L6":len(distinct),
            "known_outcome_risk_set_after_L6":len(rs),"events_after_L6_through_12m":int(sum(r["y"] for r in rs)),
            "L3_rule":next(x["L3_rule"] for x in timing if x["cohort"]==cohort),
            "outcome_rule":"exclude recurrence on/before cohort L6; recurrence through month 12 or confirmed event-free follow-up through month 12"})
    l3l6_transition_rows=[]
    transition_names=("negative_to_negative","negative_to_positive","positive_to_negative","positive_to_positive")
    for cohort,rs in l3_l6_by_cohort.items():
        for tr in transition_names:
            g=[r for r in rs if r["transition"]==tr]
            if not g: continue
            ev=int(sum(r["y"] for r in g)); lo,hi=wilson(ev,len(g))
            l3_months=[float(r["l3_day"])/DAY_PER_MONTH for r in g if r.get("l3_day") is not None]
            l6_months=[float(r["landmark_day"])/DAY_PER_MONTH for r in g]
            l3l6_transition_rows.append({"cohort":cohort,"transition":tr,"n":len(g),"events":ev,
                "recurrence_rate_after_L6_to_12m":ev/len(g),"wilson_95ci_low":lo,"wilson_95ci_high":hi,
                "L3_sample_month_median":float(np.median(l3_months)) if l3_months else None,
                "L6_origin_month_median":float(np.median(l6_months)),
                "outcome_window":"after cohort-defined L6 through postoperative month 12"})
    for tr in transition_names:
        g=[r for c,rs in l3_l6_by_cohort.items() for r in rs if r["transition"]==tr]
        if not g: continue
        ev=int(sum(r["y"] for r in g)); lo,hi=wilson(ev,len(g))
        l3l6_transition_rows.append({"cohort":"Pooled public (descriptive)","transition":tr,"n":len(g),"events":ev,
            "recurrence_rate_after_L6_to_12m":ev/len(g),"wilson_95ci_low":lo,"wilson_95ci_high":hi,
            "L3_sample_month_median":None,"L6_origin_month_median":None,
            "outcome_window":"after cohort-defined L6 through postoperative month 12; cohorts remain identified in source rows"})

    # Compare clinical-only, L3 update, and L6 update on the identical post-L6-to-month-12 risk set.
    l3l6_model_rows=[]; l3l6_increment_rows=[]
    l3l6_eval={"GALAXY":(galaxy,"OOB development estimate")}
    for cohort in ("COSMOS","ColonAiQ"):
        rs=[r for r in l3_l6_by_cohort.get(cohort,[])
            if all(r.get(k) is not None for k in ("age","sex_male","stage_high"))]
        if rs: l3l6_eval[cohort]=(rs,"frozen GALAXY model; independent cohort")
    def score_l3l6_group(cohort,rs,role,probs,y):
        labels=("Clinical-only","Clinical+L3 history","Clinical+L6 history")
        for label,pred in zip(labels,probs):
            l3l6_model_rows.append({"cohort":cohort,"analysis_role":role,"model":label,"n":len(rs),
                "events":int(y.sum()),"event_rate":float(y.mean()),**metric_values(y,pred,mod),
                "outcome_window":"after cohort-defined L6 assessment through postoperative month 12"})
        for before_i,after_i,before_name,after_name in ((0,1,"Clinical-only","Clinical+L3 history"),
                                                       (1,2,"Clinical+L3 history","Clinical+L6 history"),
                                                       (0,2,"Clinical-only","Clinical+L6 history")):
            ci=paired_delta_ci(y,probs[before_i],probs[after_i],mod,BOOT_REPS,SEED+len(cohort)+after_i)
            before=metric_values(y,probs[before_i],mod); after=metric_values(y,probs[after_i],mod)
            for metric,(low,high,nvalid) in ci.items():
                l3l6_increment_rows.append({"cohort":cohort,"n":len(rs),"events":int(y.sum()),
                    "comparison":f"{before_name} -> {after_name}","metric":metric,
                    "before":before[metric],"after":after[metric],"delta_after_minus_before":after[metric]-before[metric],
                    "bootstrap_ci_low":low,"bootstrap_ci_high":high,"valid_bootstrap_replicates":nvalid})
    galaxy_mask=np.isfinite(oob["clinical"]) & np.isfinite(oob["L3_before"]) & np.isfinite(oob["L6_after"])
    gy=np.asarray([r["y"] for r in galaxy],dtype=int)[galaxy_mask]
    gr=[r for r,keep in zip(galaxy,galaxy_mask) if keep]
    score_l3l6_group("GALAXY",gr,"paired bootstrap OOB",[oob["clinical"][galaxy_mask],oob["L3_before"][galaxy_mask],oob["L6_after"][galaxy_mask]],gy)
    pooled_ext={"rs":[],"clinical":[],"l3":[],"l6":[],"y":[]}
    for cohort,(rs,role) in l3l6_eval.items():
        if cohort=="GALAXY": continue
        y=np.asarray([r["y"] for r in rs],dtype=int)
        probs=[predict(full_models[k],rs,mod) for k in ("clinical","L3_before","L6_after")]
        score_l3l6_group(cohort,rs,role,probs,y)
        if cohort in ("COSMOS","ColonAiQ"):
            pooled_ext["rs"].extend(rs); pooled_ext["y"].extend(y.tolist())
            for key,pred in zip(("clinical","l3","l6"),probs): pooled_ext[key].extend(pred.tolist())
    if pooled_ext["rs"]:
        score_l3l6_group("Pooled external (descriptive)",pooled_ext["rs"],
            "frozen GALAXY model; COSMOS + ColonAiQ with L3 available",
            [np.asarray(pooled_ext[k]) for k in ("clinical","l3","l6")],np.asarray(pooled_ext["y"],dtype=int))

    # Same-horizon patient-level update summary, using paired OOB predictions in GALAXY.
    update_rows=[]
    for tr in ("negative_to_negative","negative_to_positive","positive_to_negative","positive_to_positive"):
        g=[r for r in galaxy if r["transition"]==tr and np.isfinite(r["risk_before_oob"]) and np.isfinite(r["risk_after_oob"])]
        if not g:
            continue
        ci=bootstrap_group_medians(g,BOOT_REPS,SEED+len(tr))
        events=sum(r["y"] for r in g); lo,hi=wilson(events,len(g))
        update_rows.append({"transition":tr,"transition_zh":{"negative_to_negative":"阴→阴","negative_to_positive":"阴→阳","positive_to_negative":"阳→阴","positive_to_positive":"阳→阳"}[tr],
            "n":len(g),"events":events,"observed_recurrence_rate":events/len(g),"event_rate_ci_low":lo,"event_rate_ci_high":hi,
            "median_risk_before":float(np.median([r["risk_before_oob"] for r in g])),
            "median_risk_after":float(np.median([r["risk_after_oob"] for r in g])),
            "median_delta_risk":float(np.median([r["risk_after_oob"]-r["risk_before_oob"] for r in g])),
            "before_ci_low":ci[0][0],"before_ci_high":ci[0][1],"after_ci_low":ci[1][0],"after_ci_high":ci[1][1],
            "delta_ci_low":ci[2][0],"delta_ci_high":ci[2][1],"prediction_method":"paired patient-level bootstrap out-of-bag",
            "bootstrap_replicates":BOOT_REPS})

    # Descriptive L3-to-L6 transitions use each cohort's designated L6 risk set.
    transition_rows=[]
    for cohort,rs in rows_by_cohort.items():
        for tr in ("negative_to_negative","negative_to_positive","positive_to_negative","positive_to_positive"):
            g=[r for r in rs if r["transition"]==tr]
            if not g: continue
            ev=sum(r["y"] for r in g); lo,hi=wilson(ev,len(g))
            transition_rows.append({"cohort":cohort,"transition":tr,"n":len(g),"events":ev,
                                    "event_rate":ev/len(g),"event_rate_ci_low":lo,"event_rate_ci_high":hi,
                                    "L6_origin_median_month":float(np.median([r.get("landmark_day",ORIGIN_DAY) for r in g])/DAY_PER_MONTH),
                                    "outcome_window":"cohort-defined L6 assessment to postoperative month 12"})

    # Molecular memory at current L6-negative status; predictions use OOB in GALAXY.
    memory_rows=[]
    memory_source = dict(validation_by_cohort)
    memory_source["GALAXY"] = rows_by_cohort["GALAXY"]
    for cohort,rs in memory_source.items():
        if not rs:
            continue
        predmap={r["patient_id"]:float(pred) for r,pred in zip(rs, predict(full_models["L6_after"],rs,mod))}
        if cohort=="GALAXY":
            predmap.update({r["patient_id"]:r["risk_after_oob"] for r in galaxy})
        for group,cond in (("never_positive",lambda r:r["ever_prior_l6"]==0),
                           ("previously_positive_now_negative",lambda r:r["ever_prior_l6"]==1)):
            g=[r for r in rs if r["l6_status"]==0 and cond(r)]
            if not g: continue
            ev=sum(r["y"] for r in g); lo,hi=wilson(ev,len(g))
            pr=[predmap[r["patient_id"]] for r in g if r["patient_id"] in predmap]
            memory_rows.append({"cohort":cohort,"memory_group":group,"n":len(g),"events":ev,
                "observed_recurrence_rate":ev/len(g),"observed_rate_ci_low":lo,"observed_rate_ci_high":hi,
                "mean_predicted_risk":float(np.mean(pr)) if pr else None,
                "median_predicted_risk":float(np.median(pr)) if pr else None,
                "prediction_source":"OOB bootstrap" if cohort=="GALAXY" else "frozen GALAXY model",
                "L6_origin_median_month":float(np.median([r.get("landmark_day",ORIGIN_DAY) for r in g])/DAY_PER_MONTH),
                "outcome_window":"cohort-defined L6 assessment to postoperative month 12"})

    # Patient-specific GALAXY update rows stay in outputs/local_only and are never pushed.
    dynamic_rows=[]; case_candidates=[]
    full_before=predict(full_models["L3_before"],galaxy,mod)
    full_after=predict(full_models["L6_after"],galaxy,mod)
    galaxy_trajectories=trajectories_by_cohort["GALAXY"]
    trajectory_before=predict(full_models["L3_before"],galaxy_trajectories,mod)
    trajectory_after=predict(full_models["L6_after"],galaxy_trajectories,mod)
    for i,r in enumerate(galaxy_trajectories):
        state=r["transition"]
        et=r.get("event_time_days")
        if et is not None:
            if et<=HORIZON_DAY: label="recurrence_by_12m"
            else: label="recurrence_after_12m"
        elif r.get("followup_days") is not None and r["followup_days"]>=HORIZON_DAY:
            label="no_recurrence_through_12m"
        else:
            label="censored_before_12m"
        d={"patient_id":r["patient_id"],"timepoint":"L6_study_window",
           "current_mrd":r["l6_status"],"ever_prior_positive":r["ever_prior_l6"],"state":state,
           "risk_before_update":float(trajectory_before[i]),"risk_after_update":float(trajectory_after[i]),
           "risk_delta":float(trajectory_after[i]-trajectory_before[i]),"recurrence_time_days":et or "",
           "recurrence_status":label,"followup_days":r.get("followup_days") or "",
           "endpoint_landmark_to_12m":r["y"] if r.get("y") is not None else "","risk_origin_day":ORIGIN_DAY}
        dynamic_rows.append(d)
        if r.get("y") is not None:
            case_candidates.append({**d,"transition_zh":{"negative_to_negative":"阴→阴","negative_to_positive":"阴→阳","positive_to_negative":"阳→阴","positive_to_positive":"阳→阳"}[state],
                                    "outcome_label":label,"color":{"negative_to_negative":"#238b67","negative_to_positive":"#e67e22","positive_to_negative":"#4c91c6","positive_to_positive":"#b82e38"}[state],
                                    "figure_path":str(local/"galaxy_l6_l12"/"digital_twin_case_examples.svg")})
    selected_cases=[]
    for tr in ("negative_to_negative","negative_to_positive","positive_to_negative","positive_to_positive"):
        g=[x for x in case_candidates if x["state"]==tr]
        if g:
            med=float(np.median([x["risk_after_update"] for x in g]))
            selected_cases.append(min(g,key=lambda x:abs(x["risk_after_update"]-med)))

    # Save aggregate files and local-only patient trajectories.
    csv_write(tables/"galaxy_model_coefficients.csv",beta_rows)
    csv_write(tables/"galaxy_internal_bootstrap_metrics.csv",boot_table)
    csv_write(tables/"external_validation_metrics.csv",metric_rows)
    csv_write(tables/"clinical_vs_mrd_increment.csv",increment_rows)
    csv_write(tables/"same_horizon_update_summary.csv",update_rows)
    csv_write(tables/"same_horizon_model_performance.csv",same_perf_rows)
    csv_write(tables/"galaxy_l6_landmark_sensitivity_comparison.csv",landmark_sensitivity_rows)
    csv_write(tables/"same_horizon_landmark_sensitivity_comparison.csv",same_horizon_sensitivity_rows)
    csv_write(tables/"galaxy_l6_l12_trajectory_main_table.csv",trajectory_main_rows)
    csv_write(tables/"molecular_memory_by_cohort.csv",memory_rows)
    csv_write(tables/"dynamic_state_transition_summary.csv",transition_rows)
    csv_write(tables/"galaxy_cohort_missingness_and_sample_audit.csv",audit)
    csv_write(tables/"mrd_window_harmonization_audit.csv",timing)
    csv_write(tables/"l3_l6_cohort_eligibility_audit.csv",l3l6_audit_rows)
    csv_write(tables/"l3_l6_l12_transition_summary.csv",l3l6_transition_rows)
    csv_write(tables/"l3_l6_l12_model_performance.csv",l3l6_model_rows)
    csv_write(tables/"l3_l6_l12_model_update_increment.csv",l3l6_increment_rows)
    csv_write(tables/"galaxy_calibration_curve_points.csv",calibration_rows)
    csv_write(local/"patient_dynamic_predictions.csv",dynamic_rows)
    csv_write(local/"digital_twin_case_examples.csv",selected_cases)
    make_figures(figs,metric_rows,update_rows,memory_rows,transition_rows,selected_cases)

    def fmt(x):
        try:
            value=float(x)
            return f"{value:.3f}" if math.isfinite(value) else "NA"
        except (ValueError,TypeError): return "NA"
    g_clin=next(r for r in metric_rows if r["cohort"]=="GALAXY" and r["model"]=="Clinical-only")
    g_mrd=next(r for r in metric_rows if r["cohort"]=="GALAXY" and r["model"]=="Clinical+MRD")
    internal_calibration_lines=[]
    for model_name in ("Clinical-only", "Clinical+MRD"):
        row={r["metric"]:r for r in boot_rows if r["model"]==model_name}
        ci,cs=row["calibration_intercept"],row["calibration_slope"]
        internal_calibration_lines.append(
            f"| {model_name} | {fmt(ci['apparent'])} | {fmt(ci['optimism_corrected'])} ({fmt(ci['corrected_ci_low'])}, {fmt(ci['corrected_ci_high'])}) | {fmt(cs['apparent'])} | {fmt(cs['optimism_corrected'])} ({fmt(cs['corrected_ci_low'])}, {fmt(cs['corrected_ci_high'])}) |")
    landmark_sensitivity_lines=[]
    for landmark in ("nominal_L6_day182.6_primary","day200_secondary_sensitivity"):
        for model_name in ("Clinical-only","Clinical+MRD"):
            a=next((r for r in landmark_sensitivity_rows if r["landmark_analysis"]==landmark and r["model"]==model_name and r["metric"]=="auc"),None)
            p=next((r for r in landmark_sensitivity_rows if r["landmark_analysis"]==landmark and r["model"]==model_name and r["metric"]=="average_precision"),None)
            b=next((r for r in landmark_sensitivity_rows if r["landmark_analysis"]==landmark and r["model"]==model_name and r["metric"]=="brier"),None)
            if a and p and b:
                label="主分析 day182.6" if landmark=="nominal_L6_day182.6_primary" else "敏感性 day200"
                landmark_sensitivity_lines.append(f"| {label} | {a['n']} | {a['events']} | {model_name} | {fmt(a['optimism_corrected'])} ({fmt(a['corrected_ci_low'])}, {fmt(a['corrected_ci_high'])}) | {fmt(p['optimism_corrected'])} ({fmt(p['corrected_ci_low'])}, {fmt(p['corrected_ci_high'])}) | {fmt(b['optimism_corrected'])} ({fmt(b['corrected_ci_low'])}, {fmt(b['corrected_ci_high'])}) |")
    ext_lines=[]
    for c in ("COSMOS","Chen","ColonAiQ","Pooled priority external"):
        a=next((r for r in metric_rows if r["cohort"]==c and r["model"]=="Clinical-only"),None)
        b=next((r for r in metric_rows if r["cohort"]==c and r["model"]=="Clinical+MRD"),None)
        if a and b:
            role = "优先独立验证" if c in ("COSMOS", "Chen") else ("探索性（事件少）" if c == "ColonAiQ" else "优先队列合并")
            l6_month = (f"{fmt(b['L6_origin_median_month'])} ({fmt(b['L6_origin_min_month'])}–{fmt(b['L6_origin_max_month'])})"
                        if b.get("L6_origin_median_month") is not None else "按各自L6节点")
            ext_lines.append(f"| {c} | {role} | {l6_month} | {b['n']} | {b['events']} | {fmt(a['auc'])} / {fmt(a['average_precision'])} / {fmt(a['brier'])} | {fmt(b['auc'])} / {fmt(b['average_precision'])} / {fmt(b['brier'])} |")
    external_calibration_lines=[]
    for c in ("COSMOS","Chen","ColonAiQ"):
        for model_name in ("Clinical-only","Clinical+MRD"):
            r=next((x for x in metric_rows if x["cohort"]==c and x["model"]==model_name),None)
            if r:
                external_calibration_lines.append(f"| {c} | {model_name} | {fmt(r['calibration_intercept'])} | {fmt(r['calibration_slope'])} |")
    group_lines=[f"| {r['transition_zh']} | {r['n']} | {r['events']} | {r['observed_recurrence_rate']:.1%} | {r['median_risk_before']:.1%} | {r['median_risk_after']:.1%} | {r['median_delta_risk']:+.1%} ({r['delta_ci_low']:+.1%}, {r['delta_ci_high']:+.1%}) |" for r in update_rows]
    perf_lines=[f"| {r['metric']} | {r['before_L3_history_oob']:.3f} | {r['after_L6_history_oob']:.3f} | {r['delta_after_minus_before']:+.3f} ({r['delta_ci_low']:+.3f}, {r['delta_ci_high']:+.3f}) |" for r in same_perf_rows]
    same_horizon_sensitivity_lines=[f"| {r['metric']} | {r['day1826_before_L3_oob']:.3f} | {r['day1826_after_L6_oob']:.3f} | {r['day1826_delta']:+.3f} | {fmt(r['day200_before_L3_oob'])} | {fmt(r['day200_after_L6_oob'])} | {fmt(r['day200_delta'])} |" for r in same_horizon_sensitivity_rows]
    memory_lines=[f"| {r['cohort']} | {r['memory_group']} | {r['n']} | {r['events']} | {r['observed_recurrence_rate']:.1%} | {r['mean_predicted_risk']:.1%} |" for r in memory_rows]
    trajectory_main_lines=[f"| {r['trajectory_zh']} | {'阳性' if r['current_l6_status']=='positive' else '阴性'} | {r['prior_mrd_history']} | {r['n']} | {r['events_6_to_12m']} | {r['recurrence_rate_6_to_12m']:.1%} ({r['recurrence_rate_wilson_95ci_low']:.1%}–{r['recurrence_rate_wilson_95ci_high']:.1%}) | {r['mean_oob_predicted_risk']:.1%} | {r['median_oob_predicted_risk']:.1%} |" for r in trajectory_main_rows]
    report=f"""# GALAXY 开发、独立公共验证与动态风险更新

## 分析目标与时间定义

开发队列为 GALAXY 2024；Chen、COSMOS 和 ColonAiQ 在模型冻结后分别验证。结局为临床复发，不含 ctDNA molecular recurrence。临床变量限定为分期、年龄、性别；GALAXY 患者级补充表没有肿瘤部位字段，因此核心模型不含 site。MRD 使用阳性/阴性，不拼接 VAF 或不同 assay 的连续值。

**主分析：名义L6 landmark。** GALAXY 原文将术后160–200天定义为6-month timepoint。本研究把该研究定义的 ctDNA 状态按名义术后6个月 day 182.6 作为 L6 输入，纳入 day 182.6 时无临床复发者，预测 **(182.6天, 365.24天]** 内临床复发。此为临床主分析。由于补充表2没有 GALAXY 患者级采血日期，day 182.6 是名义 landmark，采血实际日在160–200天内；该时间精度作为局限报告。day 200 起点的分析另列为时间窗敏感性分析。

L6按各研究自己的临床节点定义，不统一转换成同一个采血日：GALAXY采用原文160–200天研究定义窗（补充表无个体采血日，主分析风险起点按名义day182.6）；Chen按原文预设的术后6个月采样节点，在实际数据中选取4.5–7.5个月内最接近6个月的一次；COSMOS使用原始补充表中的scheduled month-6列；ColonAiQ按每3个月监测方案，选取4.5–7.5个月内最接近术后6个月的一次。Chen与ColonAiQ容差取季度采样间隔的一半，实际采样日作为该患者的风险起点。L3→L6分析也沿用各队列自己的节点：GALAXY为原文70–112天窗，COSMOS取补充表month-3列，ColonAiQ取术后约3个月样本；Chen原文术后ctDNA采样从6个月开始，没有预设L3节点，因此不进入L3→L6分析。复发结局从各自L6评估之后计算至术后12个月。研究采样方案分别见[GALAXY论文](https://doi.org/10.1038/s41591-024-03254-6)、[Chen等](https://d-nb.info/1241319898/34)、[COSMOS论文](https://pmc.ncbi.nlm.nih.gov/articles/PMC11443202/)及[ColonAiQ论文](https://pmc.ncbi.nlm.nih.gov/articles/PMC10119774/)。实际纳入人数及所选L3/L6采样月份见 `mrd_window_harmonization_audit.csv`。

## GALAXY 开发

匹配的三时间点风险集：N={len(galaxy)}，(182.6天,12月]复发={int(sum(r['y'] for r in galaxy))}。ridge logistic 的惩罚系数 λ=1；年龄以开发集均值{full_models['L6_after']['scale']['age_mean']:.2f}岁和标准差{full_models['L6_after']['scale']['age_sd']:.2f}标准化。Clinical-only 与 Clinical+MRD 在完全相同患者上比较。MRD模型变量为 L6 当前阳性及既往已观测 MRD 状态中任一次阳性；GALAXY 可用的既往列为 MRD-window/L3。外部队列按其实际 L6 前采样历史派生同一“ever observed positive”概念，检测频率不同会影响该特征的观察机会。

| 模型 | apparent AUC | optimism-corrected AUC | apparent AP | corrected AP | apparent Brier | corrected Brier |
|---|---:|---:|---:|---:|---:|---:|
| Clinical-only | {fmt(g_clin['auc'])} | {fmt(next(r['optimism_corrected'] for r in boot_rows if r['model']=='Clinical-only' and r['metric']=='auc'))} | {fmt(g_clin['average_precision'])} | {fmt(next(r['optimism_corrected'] for r in boot_rows if r['model']=='Clinical-only' and r['metric']=='average_precision'))} | {fmt(g_clin['brier'])} | {fmt(next(r['optimism_corrected'] for r in boot_rows if r['model']=='Clinical-only' and r['metric']=='brier'))} |
| Clinical+MRD | {fmt(g_mrd['auc'])} | {fmt(next(r['optimism_corrected'] for r in boot_rows if r['model']=='Clinical+MRD' and r['metric']=='auc'))} | {fmt(g_mrd['average_precision'])} | {fmt(next(r['optimism_corrected'] for r in boot_rows if r['model']=='Clinical+MRD' and r['metric']=='average_precision'))} | {fmt(g_mrd['brier'])} | {fmt(next(r['optimism_corrected'] for r in boot_rows if r['model']=='Clinical+MRD' and r['metric']=='brier'))} |

患者级 bootstrap 重抽样 {args.bootstrap} 次。校准截距、斜率及OOB校准分箱见相应表；完整回归系数与标准化参数已冻结至 `outputs/models/galaxy_frozen_ridge_models.json`。没有使用外部验证结果调参、重校准或选阈值。

### GALAXY校准

| 模型 | apparent截距 | optimism-corrected截距 (95% bootstrap CI) | apparent斜率 | optimism-corrected斜率 (95% bootstrap CI) |
|---|---:|---:|---:|---:|
{chr(10).join(internal_calibration_lines)}

### day200 次要时间窗敏感性分析

day200结果来自此前完整运行并保留在原分析目录的1000次bootstrap；它作为时间窗敏感性结果，不替代day182.6主分析。两个landmark的风险集不同，数值用于评估起点选择的稳健性。

| 分析 | N | events | 模型 | corrected AUC (95% bootstrap CI) | corrected AP (95% bootstrap CI) | corrected Brier (95% bootstrap CI) |
|---|---:|---:|---|---:|---:|---:|
{chr(10).join(landmark_sensitivity_lines)}

## 冻结模型的公共验证

表内顺序为 AUC / average precision / Brier；校准截距和斜率见 `external_validation_metrics.csv`。

| 队列 | 验证角色 | L6采样月，中位数（范围） | N | events | Clinical-only | Clinical+MRD |
|---|---|---:|---:|---:|---:|---:|
{chr(10).join(ext_lines)}

外部验证按COSMOS和Chen作为优先队列报告；ColonAiQ因可用事件少仅作探索性结果。纳入者须在各自L6评估前未复发、有可用既往MRD历史，并能判定术后12个月结局；事件要求发生在L6评估后、且不迟于术后12个月。各患者的L6到12个月预测时长会随实际评估日期略有不同，这保留了真实研究采样节奏。未做验证集特征选择、调参或重校准。MRD增量的配对bootstrap CI见 `clinical_vs_mrd_increment.csv`。样本量/事件少的队列其calibration slope与AP会不稳定，应结合N、events和区间解释。

| 队列 | 模型 | 校准截距 | 校准斜率 |
|---|---|---:|---:|
{chr(10).join(external_calibration_lines)}

## GALAXY L6→L12 主轨迹结果

以下独立主表按**当前L6状态 + L6之前已观察到的MRD历史**划分持续阴性、转阳、清除和持续阳性。事件为名义day 182.6之后至术后12个月内临床复发；“平均预测风险/中位预测风险”来自GALAXY主模型的患者级bootstrap OOB预测。既往“无阳性”指纳入模型的既往可观测结果均无阳性，不代表未检测时间段的真实阴性。

| L6→L12轨迹 | 当前L6 MRD | 既往MRD历史 | N | 6–12月复发 | 观察复发率 (95% Wilson CI) | 平均OOB预测风险 | 中位OOB预测风险 |
|---|---|---|---:|---:|---:|---:|---:|
{chr(10).join(trajectory_main_lines)}

可下载的独立汇总表为 `outputs/tables/galaxy_l6_l12_trajectory_main_table.csv`；患者级预测仍只保存在本地 `outputs/local_only/`。

## 分子记忆

在 L6 ctDNA 阴性者中，既往组定义为 L6 前至少一次已观测阳性；从未阳性组定义为所有已观测既往结果均阴性。

| 队列 | 既往状态 | N | 复发 | 观察复发率 | 平均预测风险 |
|---|---|---:|---:|---:|---:|
{chr(10).join(memory_lines)}

GALAXY 的分子记忆在 OOB 估计中有清楚分层；验证队列的“既往阳性、当前阴性”人数/事件很少，不能据此宣称已跨队列复制。

## 局限

1. GALAXY 的主分析将研究定义的160–200天状态映射到名义day 182.6 L6 landmark；实际采血日不可见，因此不能确认每个状态都在该名义日期前获得。day 200起点另作敏感性分析。
2. GALAXY 缺少患者级3/6月采血日，不能完成严格≤6月敏感性分析，也不能用真实日期拟合患者特异的动态生存模型。
3. GALAXY 没有个体 site，因此临床底模使用 stage、age、sex；无法检验加入部位后的效能。
4. 化疗时间、CEA、影像随访等没有在所有公共队列同定义，未进入公共 core model；模型估计关联与预测，不解释 MRD 清除的因果或治疗反应效应。
5. COSMOS 非复发者随访采用补充表 RFS/censor 月数；Chen/ColonAiQ 非复发者以可确认的随访采样时间判断12月结局。早期截尾或竞争事件不作为对照。

## 文件

系数、指标、增量、转移、分子记忆、缺失与时间窗审计表位于 `outputs/tables/`；指定图形位于 `outputs/figures/`。患者级轨迹及真实病例示例只写入 `outputs/local_only/`，不得提交到 GitHub。
"""
    (reports/"galaxy_development_external_validation_zh.md").write_text(report,encoding="utf-8")

    upreport=f"""# Same-horizon L3→L6 风险更新

## 目标

在同一组 GALAXY 患者中，两个模型预测完全相同的结局 **(182.6天, 365.24天] 临床复发**：before 只用临床信息及 L3/MRD-window history；after 增加 L6 ctDNA 与截至L6的既往阳性记忆。模型形式均为 ridge logistic、λ=1；同一病人风险集；组间风险使用1000次患者级 bootstrap 的 out-of-bag (OOB) 预测，以降低开发集表观拟合偏倚。

## 3月→6月状态分组

| 转移 | N | 复发 | 实际复发率 (95% Wilson CI) | before median risk | after median risk | Δrisk median (bootstrap 95% CI) |
|---|---:|---:|---:|---:|---:|---:|
{chr(10).join(group_lines)}

总体OOB性能也在相同患者和相同结局窗比较：

| 指标 | L3信息 OOB | L6信息 OOB | Δ (95% paired bootstrap CI) |
|---|---:|---:|---:|
{chr(10).join(perf_lines)}

`Δrisk = risk_after - risk_before`。负→阳应上调、阳→阴应下调的方向以估计结果为准，不预设必须成立。稀疏的转换/清除亚组需谨慎解读。完整患者级风险列只保存在 `outputs/local_only/patient_dynamic_predictions.csv`。真实代表性患者SVG也只在 `outputs/local_only/`。

同一结局窗的起点敏感性（day200为此前完成的次要分析，风险集不同）：

| 指标 | day182.6 L3 OOB | day182.6 L6 OOB | Δ | day200 L3 OOB | day200 L6 OOB | Δ |
|---|---:|---:|---:|---:|---:|---:|
{chr(10).join(same_horizon_sensitivity_lines)}

## 解释

这是固定未来结局窗口上的纵向风险更新，不是比较 L3→12 与 L6→12 两个不同长度预测窗。它支持回答“同一患者在看到新L6 MRD后，面对同一结局窗口的风险估计如何变化”。由于两模型均在GALAXY开发，OOB预测用于组别汇总，独立可迁移性由冻结模型的Chen/COSMOS/ColonAiQ验证另行判断。
"""
    (reports/"same_horizon_dynamic_update_zh.md").write_text(upreport,encoding="utf-8")

    dynreport=f"""# Dynamic survival sensitivity 与时间泄漏审计

本轮完成了固定结局窗口的 same-horizon 更新，不将其称作连续时间 joint model 或 time-dependent Cox。

GALAXY 原文将术后160–200天称为6-month timepoint。本轮主分析按用户指定，以名义day 182.6作为 L6 landmark，预测至术后365.24天。补充表没有患者级3/6月采血日期，不能把个体时间变 covariate 放在确切采血时点，也不能检验采血与复发的逐患者先后顺序。因此本轮**不拟合 time-dependent Cox 或 landmark supermodel**，不插值采血日。day 200 作为次要敏感性分析，用于检验等待完整L6窗口结束的影响。

## 主要泄漏检查

- GALAXY主分析采用研究定义的160–200天L6状态和名义day182.6风险起点；每个外部队列按自己的study-defined/clinically designated L6节点使用对应结果，没有跨队列统一的日数截止。
- 各队列分别排除L6评估前或当日已复发者；预测结局是L6评估后至术后365.24天的临床复发。GALAXY个体采血日不可见，故只能使用名义day182.6；该研究窗内的实际事件与采血先后关系不能逐患者核验。
- 复发晚于术后365.24天只作为12月前无复发对照；非复发对照须有至少365.24天可确认随访。
- 未观测 ctDNA 不编码为阴性；临床结局早期截尾或随访不足者排除于拟合/验证集。
- 外部队列未用于特征选择、惩罚参数选择、阈值选择或重校准。
- 患者级同一结局窗风险更新用同一人群与配对OOB风险，避免将不同预测时距的风险直接相减。

## 留待未来院内数据

要估计真正患者特异的 `P(T_recurrence > u | MRD history through t)`，需院内队列同时保留确切 surgery date、每次采血日期、复发日期及末次临床无复发评估日期；届时优先用 landmark supermodel 或 time-dependent Cox，并报告 time-dependent AUC、C-index、Brier/IBS 与校准。当前公开GALAXY表的时间精度不支持这些指标，因此本轮不报告。
"""
    (reports/"dynamic_model_sensitivity_zh.md").write_text(dynreport,encoding="utf-8")

    transition_zh={"negative_to_negative":"L3阴→L6阴","negative_to_positive":"L3阴→L6阳",
                   "positive_to_negative":"L3阳→L6阴","positive_to_positive":"L3阳→L6阳"}
    l3l6_audit_lines=[f"| {r['cohort']} | {r['L3_rule']} | {r['L3_assessments']} | {r['L6_assessments']} | {r['patients_with_both_statuses']} | {r['known_outcome_risk_set_after_L6']} | {r['events_after_L6_through_12m']} |"
                      for r in l3l6_audit_rows]
    l3l6_transition_lines=[f"| {r['cohort']} | {transition_zh[r['transition']]} | {r['n']} | {r['events']} | {r['recurrence_rate_after_L6_to_12m']:.1%} ({r['wilson_95ci_low']:.1%}–{r['wilson_95ci_high']:.1%}) | {fmt(r['L3_sample_month_median']) if r['L3_sample_month_median'] is not None else '研究定义窗'} | {fmt(r['L6_origin_month_median']) if r['L6_origin_month_median'] is not None else '队列特异'} |"
                            for r in l3l6_transition_rows]
    l3l6_model_lines=[f"| {r['cohort']} | {r['model']} | {r['n']} | {r['events']} | {fmt(r['auc'])} | {fmt(r['average_precision'])} | {fmt(r['brier'])} |"
                      for r in l3l6_model_rows]
    l3l6_increment_lines=[f"| {r['cohort']} | {r['comparison']} | {r['metric']} | {r['delta_after_minus_before']:+.3f} ({fmt(r['bootstrap_ci_low'])}, {fmt(r['bootstrap_ci_high'])}) |"
                          for r in l3l6_increment_rows]
    l3l6_report=f"""# 队列特异节点的 L3→L6→L12 动态分析

## 时间点定义

- **GALAXY**：L3采用论文定义的术后70–112天窗口；L6采用160–200天窗口，因补充表无患者级采血日期，L6→L12结局风险起点仍按预先确定的day182.6。
- **COSMOS**：L3和L6分别采用补充表中scheduled month-3和month-6节点；结局起点为患者的L6评估。
- **ColonAiQ**：按季度监测记录选择最接近术后3个月、且早于L6的实际样本作为L3；L6按最近6个月节点选择。每个患者的实际L6样本日作为结局起点。
- **Chen**：原文术后ctDNA采样在day3–7后从术后6个月开始，此后每3个月；没有研究定义的3个月MRD节点，因此不纳入L3→L6分析。少数非计划早期采样不重新标作L3。

各队列结局为：患者在其L6评估前无复发，并统计L6评估之后至术后12个月的临床复发。只有能够确认12个月结局者进入轨迹率与模型评价；未知ctDNA不编码为阴性。采样方案参见[GALAXY论文](https://doi.org/10.1038/s41591-024-03254-6)、[Chen队列原文](https://pmc.ncbi.nlm.nih.gov/articles/PMC8130394/)、[COSMOS论文](https://pmc.ncbi.nlm.nih.gov/articles/PMC11443202/)和[ColonAiQ论文](https://pmc.ncbi.nlm.nih.gov/articles/PMC10119774/)。

## L3/L6可用性与12个月结局风险集

| 队列 | L3节点 | 有L3 | 有L6 | 两节点均有 | L6后结局可判定 | 其中复发 |
|---|---|---:|---:|---:|---:|---:|
{chr(10).join(l3l6_audit_lines)}

## L3→L6转移与L6后复发

| 队列 | 轨迹 | N | L6后至12月复发 | 复发率 (95% Wilson CI) | L3采样月中位数 | L6起点月中位数 |
|---|---|---:|---:|---:|---:|---:|
{chr(10).join(l3l6_transition_lines)}

合并行只作描述性汇总；具体分析仍以各队列分层结果为主。不同队列assay、病例构成和采样节奏不同，不能把该合并比例解释成单一目标人群风险。

## same-horizon 预测更新

Clinical-only、Clinical+L3 history和Clinical+L6 history均在相同的L6后至12个月风险集上比较。GALAXY报告患者级bootstrap的OOB预测；外部队列使用冻结的GALAXY系数，不重拟合、不校准。Chen因缺少L3节点不参与这组比较。

| 队列 | 模型 | N | 事件 | AUC | AP | Brier |
|---|---|---:|---:|---:|---:|---:|
{chr(10).join(l3l6_model_lines)}

| 队列 | 更新比较 | 指标 | 更新后−更新前 (95% paired bootstrap CI) |
|---|---|---|---:|
{chr(10).join(l3l6_increment_lines)}

上述同一结局窗比较检验在L3获得的信息基础上，加入L6状态是否进一步更新风险排序与误差。外部队列事件数有限，尤其ColonAiQ，AUC/AP及其差值区间应谨慎解读。L6 landmark前复发者被排除，因此本分析回答的是“L6时仍未复发者未来至12个月的风险”，不代表从L3起至12个月的总复发风险。

详细数据见 `outputs/tables/l3_l6_cohort_eligibility_audit.csv`、`l3_l6_l12_transition_summary.csv`、`l3_l6_l12_model_performance.csv` 与 `l3_l6_l12_model_update_increment.csv`。患者级预测文件仅保存于本地 `outputs/local_only/`。
"""
    (reports/"l3_l6_l12_dynamic_summary_zh.md").write_text(l3l6_report,encoding="utf-8")

    # Compact stdout enables a quick, independent review of denominators after the run.
    print(json.dumps({"origin_day":ORIGIN_DAY,"development_n":len(galaxy),"development_events":int(sum(r["y"] for r in galaxy)),
        "transition_complete_cohorts":{c:{"n":len(rs),"events":int(sum(r["y"] for r in rs))} for c,rs in rows_by_cohort.items()},
        "validation_cohorts":{c:{"n":len(rs),"events":int(sum(r["y"] for r in rs))} for c,rs in validation_by_cohort.items()},
        "updates":update_rows,"memory":memory_rows,"bootstrap_successes":args.bootstrap,
        "site_available_in_galaxy":False},ensure_ascii=False))


if __name__ == "__main__":
    main()
